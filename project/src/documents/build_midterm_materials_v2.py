"""Evidence-bound course midterm revision, stage summary and editable briefing slides."""
from pathlib import Path
from copy import deepcopy
from hashlib import sha256
import json,sys,re
if __package__ in (None,''):sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
import numpy as np
import matplotlib
matplotlib.use('Agg')
from matplotlib import pyplot as plt
from matplotlib.patches import FancyBboxPatch,FancyArrowPatch
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from pptx import Presentation
from pptx.util import Inches,Pt as PPt
from pptx.dml.color import RGBColor as PRGB
from pptx.enum.shapes import MSO_SHAPE,MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR
from documents.build_midterm_report import clean,p,table,font,unsnap

ROOT=Path(__file__).resolve().parents[3];OUT=ROOT/'中期报告相关'
RECORD=ROOT/'项目导航/整理记录_2026-10-02_v1';ASSETS=RECORD/'审阅'
TITLE='基于主动探索的视觉地理定位方法设计与实现'
REPORT=OUT/f'中期报告_{TITLE}.docx'
SUMMARY=OUT/'阶段提升总结_v2.docx'
SLIDES=OUT/f'答辩材料/中期汇报_{TITLE}.pptx'
PROPOSAL=ROOT/f'选题报告相关/最新版选题报告/18-{TITLE}-陈德昕-李世显-陈龙辉.docx'


def read(p):return json.loads(p.read_text('utf8'))
def digest(p):return sha256(p.read_bytes()).hexdigest()
def new_json(p,x):
    with p.open('x',encoding='utf8') as f:json.dump(x,f,ensure_ascii=False,indent=2)


def evidence():
    folders=[('S2','Masa','边缘线索S2正式复验_v1',100,4,5),('S3','Masa','边缘线索S3独立源图确认_v1',250,10,5),
        ('S4','Masa','S4十乘十正式扩展_v1',500,20,10),('Swiss','SwissView','五乘五正式迁移_v1',500,20,5)]
    e={};paths={}
    for code,dataset,name,tasks,sources,k in folders:
        folder=ROOT/f'DATA/processed_data/{dataset}/评测结果/{name}';s=read(folder/'对照汇总.json');v=read(folder/'验收结论.json')
        if not any(v.get(key) is True for key in ('formal_S2_passed','formal_S3_passed','formal_S4_passed','formal_cross_dataset_passed')):raise ValueError('accepted main method required')
        if digest(folder/'独立复核.json')!=v['audit_sha256'] or digest(folder/'对照汇总.json')!=v['summary_sha256']:raise ValueError('formal binding drift')
        e[code]=dict(**s['averages']['Edge']['CueFull'],tasks=tasks,sources=sources,k=k,budget=2*k,path=folder.relative_to(ROOT).as_posix(),
            baseline=s['averages']['Edge']['Baseline'],mean=s['averages']['Edge']['CueMean'],wrong=s['averages']['Edge']['CueWrong'])
        for n in ('对照汇总.json','验收结论.json','独立复核.json'):paths[(folder/n).relative_to(ROOT).as_posix()]=digest(folder/n)
    for code,dataset,name,summary in [('Target','Masa','继续训练冻结目标证据复验_v1','目标证据汇总.json'),
        ('Fresh','SwissView','继续训练独立源图确认_v1','主对照汇总.json'),('Protocol','Masa','探索器协议适配开发验证_v1','主对照/对照汇总.json')]:
        folder=ROOT/f'DATA/processed_data/{dataset}/评测结果/{name}';v=read(folder/'验收结论.json')
        if not v['audit_passed'] or digest(folder/'独立复核.json')!=v['audit_sha256'] or digest(folder/summary)!=v['summary_sha256']:raise ValueError('new candidate audit binding')
        e[code]=read(folder/summary);e[code]['verdict']=v;e[code]['path']=folder.relative_to(ROOT).as_posix()
        for n in (summary,'验收结论.json','独立复核.json'):paths[(folder/n).relative_to(ROOT).as_posix()]=digest(folder/n)
    default=ROOT/'project/local_policy_default.json';assert digest(default)=='bb944dc9acf5ec0760d64e83a7599d9a61c1e4a8c736d3d9cfabbe4afcde874b'
    paths[default.relative_to(ROOT).as_posix()]=digest(default)
    for n in ['选题报告相关/任务要求与论文框架对应说明_v1.md','选题报告相关/SwissView源文件使用状态_2026-10-02_v1.json',
        'DATA/processed_data/Masa/评测结果/证据账本简化接口验证_v2/阶段G/对照汇总.json']:
        pth=ROOT/n
        if pth.is_file():paths[n]=digest(pth)
    paths[PROPOSAL.relative_to(ROOT).as_posix()]=digest(PROPOSAL)
    paths['课程相关报告模版/综合工程设计-2-中期报告.doc']=digest(ROOT/'课程相关报告模版/综合工程设计-2-中期报告.doc')
    return e,paths


def figures(e):
    plt.rcParams.update({'font.family':'Microsoft YaHei','font.size':10,'axes.labelsize':10,'axes.titlesize':11,
        'xtick.labelsize':9,'ytick.labelsize':9,'axes.spines.top':False,'axes.spines.right':False,
        'pdf.fonttype':42,'svg.fonttype':'none','axes.unicode_minus':False})
    f,ax=plt.subplots(figsize=(5.4,2.45));names=['S2稳定\n5×5','S3独立源图\n5×5','网格密度\n10×10','SwissView迁移\n5×5']
    vals=[e[c]['sr_mean']*100 for c in ('S2','S3','S4','Swiss')];cols=['#356B54','#356B54','#B57B30','#477A82']
    ax.bar(range(4),vals,color=cols,width=.58,zorder=2)
    for i,v in enumerate(vals):ax.text(i,v+1.6,f'{v:.2f}%',ha='center',fontsize=10)
    ax.set_xticks(range(4),names);ax.set_ylim(0,105);ax.set_yticks([0,20,40,60,80,100]);ax.set_ylabel('SR（%）');ax.grid(axis='y',color='#DEE4DF',lw=.6);ax.set_axisbelow(True)
    f.subplots_adjust(left=.12,right=.98,bottom=.28,top=.97)
    f.savefig(ASSETS/'报告_主线验收.png',dpi=300);f.savefig(ASSETS/'报告_主线验收.pdf');plt.close(f)
    f,axes=plt.subplots(1,2,figsize=(5.4,2.65));a=e['Fresh']['averages']
    for j,(metric,scale,label) in enumerate([('sr',100,'SR（%）'),('sg',1,'SG（格）')]):
        ax=axes[j]
        for arm,col,off in [('M0','#87918C',-.17),('Continue5','#356B54',.17)]:
            vs=[a[k][arm][metric+'_mean']*scale for k in ('5','10')]
            ax.bar(np.array([0,1])+off,vs,width=.30,color=col,label=arm,zorder=2)
            for i,v in enumerate(vs):ax.text(i+off,v+(1 if j==0 else .035),f'{v:.2f}' if j==0 else f'{v:.3f}',ha='center',fontsize=9)
        ax.set_xticks([0,1],['5×5','10×10']);ax.set_ylabel(label);ax.grid(axis='y',color='#DEE4DF',lw=.6);ax.set_axisbelow(True)
        ax.set_ylim(0,110 if j==0 else 1.65)
    axes[0].legend(loc='upper center',ncol=2,fontsize=8.5,frameon=False)
    f.subplots_adjust(left=.1,right=.98,bottom=.22,top=.97,wspace=.50)
    f.savefig(ASSETS/'报告_新源文件确认.png',dpi=300);f.savefig(ASSETS/'报告_新源文件确认.pdf');plt.close(f)
    # Presentation derivative: larger labels, immutable scientific originals stay intact.
    f,axes=plt.subplots(1,2,figsize=(9.3,4.45))
    arms=e['Target']['arms'];keys=['Baseline','CueFull','CueMean','CueWrong']
    labels=['禁用线索','真实目标','均值目标','错误目标']
    for ax,metric,scale,label,ceiling in zip(axes,['sr_mean','sg_mean'],[100,1],['SR（%）','SG（格）'],[105,2.3]):
        vs=[arms[k][metric]*scale for k in keys]
        ax.bar(range(4),vs,color=['#87918C','#356B54','#B57B30','#477A82'],width=.63)
        ax.set_xticks(range(4),labels,rotation=20,ha='right',fontsize=15);ax.tick_params(axis='y',labelsize=13)
        ax.set_ylabel(label,fontsize=16);ax.set_ylim(0,ceiling);ax.grid(axis='y',color='#DEE4DF',lw=.8);ax.set_axisbelow(True)
        for i,v in enumerate(vs):ax.text(i,v+(2.5 if scale==100 else .065),f'{v:.2f}' if scale==100 else f'{v:.3f}',ha='center',fontsize=15)
    f.subplots_adjust(left=.09,right=.98,bottom=.24,top=.92,wspace=.46)
    f.savefig(ASSETS/'汇报_目标证据简图.png',dpi=300);f.savefig(ASSETS/'汇报_目标证据简图.pdf');plt.close(f)
    f,ax=plt.subplots(figsize=(5.4,2.75));ax.set_xlim(0,10);ax.set_ylim(0,5.0);ax.axis('off')
    boxes=[(.15,3.45,2.0,.92,'目标图与当前图\n公开访问状态','#F0F4F0'),(2.75,3.45,2.1,.92,'冻结 Sat2Cap\n图像表示','#E7EFEB'),
        (5.45,3.60,2.2,.72,'Small256 探索器','#E7EFEB'),(5.45,2.3,2.2,.92,'EdgeTargetCue\n目标与当前图匹配','#E7EFEB'),
        (7.95,3.45,1.8,.92,'可信线索门控\n合法动作筛选','#E7EFEB'),(7.95,.85,1.8,.92,'网格环境\n移动与反馈','#F0F4F0')]
    for x,y,w,h,t,col in boxes:
        ax.add_patch(FancyBboxPatch((x,y),w,h,boxstyle='round,pad=.03,rounding_size=.08',facecolor=col,edgecolor='#456B55',lw=.8))
        ax.text(x+w/2,y+h/2,t,ha='center',va='center',fontsize=9)
    for start,end in [((2.18,3.9),(2.70,3.9)),((4.9,3.9),(5.40,3.9)),((4.9,3.5),(5.42,2.95)),((7.7,3.9),(7.9,3.9)),((7.7,2.85),(8.3,3.4)),((8.85,3.4),(8.85,1.8)),((7.9,1.3),(1.2,1.3)),((1.2,1.3),(1.2,3.4))]:
        ax.add_patch(FancyArrowPatch(start,end,arrowstyle='->',mutation_scale=10,color='#456B55',lw=.9))
    ax.text(4.8,.83,'连续局部观察与行动反馈',ha='center',fontsize=9)
    ax.text(5,.12,'目标坐标、真实距离及未访问图块不进入执行策略',ha='center',fontsize=9)
    f.subplots_adjust(left=.02,right=.98,top=.98,bottom=.03)
    f.savefig(ASSETS/'报告_Agent闭环.png',dpi=300);f.savefig(ASSETS/'报告_Agent闭环.pdf');plt.close(f)


def picture(cell,name,width=13.5):
    para=cell.add_paragraph();unsnap(para);para.alignment=WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_before=Pt(2);para.paragraph_format.space_after=Pt(3)
    run=para.add_run();run.add_picture(str(ASSETS/name),width=Cm(width))
    para.paragraph_format.keep_with_next=True


def append_page(d,anchor,prototype,content):
    sep=OxmlElement('w:p');pr=OxmlElement('w:pPr');flag=OxmlElement('w:pageBreakBefore');pr.append(flag);sep.append(pr)
    anchor.addnext(sep);t=OxmlElement('w:tbl');t.append(deepcopy(d.tables[0]._tbl.tblPr));t.append(deepcopy(d.tables[0]._tbl.tblGrid))
    tr=deepcopy(prototype);tc=tr.find(qn('w:tc'))
    for ch in list(tc):
        if ch.tag!=qn('w:tcPr'):tc.remove(ch)
    tc.append(OxmlElement('w:p'));t.append(tr);sep.addnext(t)
    from docx.table import _Cell
    content(_Cell(tc,d.tables[0]));return t


def source_body(e):
    return [
        ('E1','S2稳定验证','Masa/评测结果/边缘线索S2正式复验_v1'),
        ('E2','S3独立源文件确认','Masa/评测结果/边缘线索S3独立源图确认_v1'),
        ('E3','S4十乘十密度扩展','Masa/评测结果/S4十乘十正式扩展_v1'),
        ('E4','SwissView同模态迁移','SwissView/评测结果/五乘五正式迁移_v1'),
        ('E5','冻结目标图利用证据','Masa/评测结果/继续训练冻结目标证据复验_v1'),
        ('E6','继续训练独立源图确认','SwissView/评测结果/继续训练独立源图确认_v1'),
        ('E7','同预算训练协议对照','Masa/评测结果/探索器协议适配开发验证_v1'),
        ('E8','同预算多角色闭环','Masa/评测结果/证据账本简化接口验证_v2')]


def build_report(e):
    d=Document(OUT/'历史版本/2026-09-30_中期报告_修订前.docx');big=d.tables[0];prototype=deepcopy(d.tables[1].rows[0]._tr)
    # Keep the cover/metadata, teacher grid and page geometry; replace historical narrative.
    b=d._element.body;start=list(b).index(big._tbl)
    for child in list(b)[start+1:]:
        if child.tag!=qn('w:sectPr'):b.remove(child)
    proposal=Document(PROPOSAL);personal=proposal.tables[0]
    for ri in range(2,6):
        src=personal.rows[ri].cells;dst=big.rows[ri].cells;seen=set()
        for ci,c in enumerate(dst):
            if c._tc in seen:continue
            seen.add(c._tc);text=src[ci].text
            if text.strip() and text not in ('申请人(小组)','完成人(小组)'):
                clean(c);para=c.paragraphs[0];para.alignment=WD_ALIGN_PARAGRAPH.CENTER;para.paragraph_format.line_spacing=1.15;font(para.add_run(text),10.5)
    for para in d.paragraphs:
        t=para.text.strip()
        if t.startswith('小组人员'):para.clear();font(para.add_run('小组人员：陈德昕  李世显  陈龙辉'),15)
        elif t.startswith('指导老师'):para.clear();font(para.add_run('指导老师：王笛'),15)
        elif t.startswith('填写时间'):para.clear();font(para.add_run('填写时间：2026 年 10 月 2 日'),15)
    c=big.rows[6].cells[0];clean(c)
    p(c,'一、设计目标与任务','title')
    p(c,'本项目面向无人机等智能体的主动视觉地理定位任务，构建局部航拍观察与环境交互平台。给定目标图像和起始位置，智能体利用连续获得的视觉信息，自主选择探索方向，在有限行动预算内到达目标网格。当前已完成平台、单智能体主方法及分阶段验收，并完成网格密度与跨数据集同模态扩展。')
    p(c,'任务采用上、右、下、左四方向移动。原协议为5×5网格、10步预算，扩展协议为10×10网格、20步预算。越界移动停留原地并扣一步；首次到达目标立即成功，最后一步到达仍计成功。策略可读取目标图、当前图、公开位置、剩余预算和访问历史；目标坐标、真实距离和未访问图块仅由环境及评测端掌握。')
    p(c,'本地方法借鉴GOMAA-Geo的主动定位任务、航拍切图与序列决策思路，采用小型GRU和强化学习适配本机资源；GeoExplorer的好奇心探索与DynCur-Geo的动态奖励思路已作受控对照。上述借鉴不等于完整复现原论文，多模态目标与全部论文基准尚未覆盖。')
    p(c,'当前正式要求允许训练小模型，多Agent属于可选拓展。已经验收的主方法是训练过的单导航智能体模块系统。后续以成功率、目标图贡献、同预算比较和独立源文件证据共同评价方法，不以增加角色或模型规模本身作为进展。')
    pages=[]
    def platform(c):
        p(c,'二、完成情况','title');p(c,'（一）平台 数据与统一评测','sub')
        p(c,'已完成图块处理、任务清单、局部观察环境、四方向动作、奖励构造及终止规则修复；训练与推理保持合法动作筛选一致。训练可用真实距离构造奖励，执行策略不读取该真值。视觉输入去除文件名和元数据线索，源文件与图块哈希记录来源。')
        table(c,[['数据来源','本地源文件','当前角色'],['Masa航拍','137训练 / 4验证 / 10测试','探索器109拟合；28已知开发留出'],['视觉头训练划分','87拟合 / 22校准','均值与0.50阈值先冻结'],['SwissView100','100张；48已用于导航','跨数据集同模态验证；52尚未跑导航']], [4.3,4.0,5.6])
        p(c,'Masa来自Massachusetts Buildings公开航拍资料，当前本地规模为151源文件；SwissView100为瑞士航拍子集。源文件隔离不表示严格地理足迹无重叠，也不排除冻结预训练编码器曾覆盖相同地域。任务ID、重复路线与独立源文件分开统计。')
        p(c,'SR为预算内成功题数除以计划题数；SG为全部终局的曼哈顿距离均值，成功记0。另记录工程完成率、重访、越界、分距离指标和调用时延。云端中断单列并按协议处理，不能用程序完成替代定位成功。')
        p(c,'每批保留预登记、源名单、模型与数据哈希、源码快照、逐步动作、终局指标和重放审计。三份训练权重全部纳入，API重复不当训练种子，重复题不增加独立地图。主线各批验收通过，最新两批合计12000条轨迹、181282动作完成重放。')
    pages.append(platform)
    def method(c):
        p(c,'（二）单智能体方法与目标视觉证据','sub')
        p(c,'主方法采用冻结Sat2Cap图像编码器、约80万参数Small256 GRU探索器和EdgeTargetCue视觉线索头。探索器以当前图像、跨步状态和公开访问信息产生探索动作；NoTarget表示探索器的目标通道使用固定训练均值，系统中的目标匹配由专门视觉头完成。')
        p(c,'视觉头比较给定目标与当前图的语义、局部关系及细边缘连续性，判断可信邻接方向。仅当最高方向置信度达到0.50、动作合法且落点未访问时接管；其余状态沿用探索器。策略每步只获得当前图与给定目标图对应特征，不查询未知邻格图像。')
        picture(c,'报告_Agent闭环.png');p(c,'图1  当前单智能体的观察 决策 行动反馈闭环','caption')
        p(c,'S2固定val100×3同题对照中，无视觉线索SR为73.00%，真实目标为89.33%，均值与错误目标为69.00%与69.67%。目标贡献与工程门槛均通过，说明本队列上目标线索能够增加实际成功，而非仅改变输出动作。','note')
        p(c,'PBRS与合法动作筛选作为基础训练与执行机制保留；好奇心奖励、网络扩容和若干目标监督候选未形成稳定收益。探索器与视觉头是单智能体内部模块，不能将这一结构称为多智能体协作。')
    pages.append(method)
    def formal(c):
        p(c,'（三）正式分阶段验收与迁移结果','sub')
        rows=[['验收','源图 / 每权重任务','协议','平均SR','平均SG']]
        for code,label in [('S2','S2稳定验证'),('S3','S3独立源图'),('S4','S4密度扩展'),('Swiss','SwissView迁移')]:
            x=e[code];rows.append([label,f"{x['sources']} / {x['tasks']}",f"{x['k']}×{x['k']} / B{x['budget']}",f"{x['sr_mean']:.2%}",f"{x['sg_mean']:.3f}"])
        p(c,'表1  当前冻结主方法的已通过验收','caption');table(c,rows,[3.3,3.1,2.6,2.5,2.4])
        p(c,'表内均为三份冻结训练权重的平均结果。S2使用4验证源文件，S3使用10测试源文件；S4的20源文件已用于此前开发，属于密度扩展工程确认。SwissView迁移使用固定20新源文件，未为迁移进行任务适配训练。','note')
        picture(c,'报告_主线验收.png');p(c,'图2  四项阶段验收协议的成功率','caption')
        p(c,'各队列的数据、距离和预算不同，图表用于分别展示验收结果，不能把不同队列的分数连接为受控提升曲线，也不能直接与论文不同协议的结果相减。10×10仍在同一原图足迹内提高切图密度，尚未扩大实际地理覆盖范围。')
    pages.append(formal)
    def recent(c):
        p(c,'（四）继续训练候选与独立确认','sub')
        p(c,'最近保持Small256、目标头、0.50阈值和合法动作规则，比较5×5继续训练Continue5与10×10适配Adapt10。两臂各三模型、每模型81920实际训练动作与1024优化器更新，全部训练及复核结束后统一使用最终权重评测，未挑最好轮次或种子。')
        p(c,'同20已见Masa源文件、500题×3的10×10评测：原M0为77.13%/SG1.201，Continue5为79.73%/0.823，Adapt10为80.00%/1.059。Adapt10较Continue5仅增加0.27点且SG更差，未通过同预算协议升级条件。')
        a=e['Target']['arms'];p(c,'表2  Continue5的冻结目标对照','caption')
        table(c,[['条件','平均SR','平均SG'],*[ [label,f"{a[key]['sr_mean']:.2%}",f"{a[key]['sg_mean']:.3f}"] for key,label in [('CueFull','真实目标'),('Baseline','禁用视觉线索'),('CueMean','均值目标'),('CueWrong','错误目标')]]],[7.1,3.4,3.4])
        p(c,'真实目标1500条按SHA复用并重新审计，三控制新增4500条。均值与错目标同时作用于全局、局部、边缘目标通道；环境仍按真实目标判分。真实目标对三控制均超过预定5点SR收益、至少两份权重为正且SG不差，目标证据通过。','note')
        p(c,'随后固定此前预留SwissView40—59号20新源文件，分别执行5×5/B10与10×10/B20，原M0与Continue5共6000条主轨迹。任务和源名单在新导航前冻结，不因结果更换样本。该20源文件已消费，后续不能重复当作未见确认。')
    pages.append(recent)
    def independent(c):
        p(c,'（四）继续训练候选与独立确认 续','sub')
        a=e['Fresh']['averages'];p(c,'表3  同20新源文件的逐题配对结果','caption')
        table(c,[['新源协议','原M0 SR / SG','Continue5 SR / SG'],*[ [f'{k}×{k}',f"{a[k]['M0']['sr_mean']:.2%} / {a[k]['M0']['sg_mean']:.3f}",f"{a[k]['Continue5']['sr_mean']:.2%} / {a[k]['Continue5']['sg_mean']:.3f}"] for k in ('5','10') ]],[4.3,4.8,4.8])
        picture(c,'报告_新源文件确认.png');p(c,'图3  独立新源文件的成功率与全终局距离','caption')
        p(c,'5×5达到预定性能保护要求。新10×10的SR仅增加0.60个百分点，三权重为持平、下降1.60点、提升3.40点，只有1/3正收益；按20源文件成组重采样的95% SR差区间为[−0.87，+2.27]点。SG下降0.362格，但不能替代成功率和一致性要求。因此新候选未通过升级，原默认及已验收主线保持。')
        p(c,'新10×10配对恢复185条、损伤176条，净增加9/1500条；5×5恢复68、损伤58。保留双方的恢复与损伤，避免只展示新增成功。主门槛失败后未追加4500条新图目标对照，也未调参数或继续训练追求通过。')
        p(c,'6000新源轨迹、74987动作重放一致；2500派生图块从原图重建，编码器与边缘特征重新计算一致。全部3000条5×5记录与原推理包装器逐步相同，排除包装器改动造成的原协议行为差异。','note')
    pages.append(independent)
    def factors(c):
        p(c,'（五）其他受控探索与阶段产出','sub')
        table(c,[['因素','已完成验证','处理结果'],['目标视觉匹配','语义、局部关系、邻接判定、细边缘线索','边缘线索形成已验收主方案'],['旋转与接受可靠性','四朝向搜索、阈值校准、旋转负样本与高分负例','存在精度/召回权衡，候选未稳定升级'],['多角色协作','同模型、同信息和调用预算的反思/双角色对照','接口修复后仍无足够SR收益，支线收束'],['零调用账本规划','完整500题×3及保护/邻域覆盖单因素','账本SR+1.27点，低成本备选；保护因素未过'],['训练协议适配','Small256同预算5×5继续训练与10×10适配','适配未过；继续训练目标证据过、独立升级未过']],[3.1,5.5,5.3])
        p(c,'云端Gemma与DeepSeek已接入并完成同策略小队列比较；接口成功与策略有效分开记录。云端完整S2曾受超时、限流影响，不能把本地正式验收移用于云端。后续协作先完成静态触发、消息诊断和同预算闭环，已测配置未支持扩大Agent数量。')
        p(c,'当前材料包括完整实验索引、冻结配置、训练/评测代码、逐步日志、哈希与复核回执；科研目录已有35张结果图、2张架构/边界图，包含PNG、PDF和SVG来源清单。报告、阶段总结及中期汇报采用同一指标口径，失败候选按原批次保留。')
        p(c,'本次材料优化不新增训练或导航评测。原始数据、模型、代码和实验路径保持；历史说明按当时状态归档，新结果通过当前入口索引。重复快照承载原批证据，不因字节相同而删除。')
    pages.append(factors)
    def issues(c):
        p(c,'三、存在的主要问题','title')
        p(c,'1. 适用条件仍有限。当前主方案在同源、连续、方向一致的航拍切图中利用边缘连续性。真实异时、季节、任意角度、跨视角或不同目标模态尚未验证；该线索可能随成像变化失效。已完成的目标扰动实验用于界定这一范围，不据干净数据高SR推断真实无人机鲁棒性。')
        p(c,'2. 更大实际区域仍待验证。10×10是同一原图内更密集的切分，并非扩大真实覆盖面积。当前定位输出为目标网格，尚未输出经纬度或接入飞控、里程计和真实航行安全约束。大范围未知区域的课程问题已在仿真中建立，实际部署仍有工作。')
        p(c,'3. 独立样本与统计证据有限。S2仅4源文件，任务可能重复；SwissView累计48文件用于导航，52尚未使用。最新双网格共用20文件，不能称40独立地图。按源文件成组计算区间仍不证明地理无重叠或预训练编码器无覆盖。')
        p(c,'4. 新因素收益不稳定。Continue5在已见开发队列上SR增加2.60点，在新10×10队列只增加0.60点。相同题上恢复185条也损伤176条，说明策略替换存在权衡；SG降低具有意义，但稳定增加预算内成功仍是首要问题。')
        p(c,'5. 多Agent与云端能力属于未证实拓展。当前主系统仍是单智能体模块组合，已有多角色小实验没有稳定收益。云端输出和视觉转发问题需按服务配置分别验证；本地验收不代表任意云端模型有效。正式任务允许训练，整体系统不能称零训练。')
        p(c,'6. 文档与工程维护需要持续同步。早期报告停留在66%与S2未通过，已与当前主线不一致。本次修订将阶段判定、样本数量和失效范围统一；历史版本保留原状态，后续新实验须另建批次，不覆盖旧验收。')
    pages.append(issues)
    def next_steps(c):
        p(c,'四、下一步需要完成的工作','title')
        p(c,'先收束已完成探索，保留当前已验收默认及Continue5的低SG结果，完善课程中期检查和答辩材料。新增实验应由明确失败证据驱动，采用单因素、固定预算和独立确认；不持续扩大网络、角色数量或重复使用同一开发集调到过线。')
        table(c,[['顺序','后续工作','检查点与交付'],['近期','同步中期报告、阶段总结、汇报与目录索引','指标来源、样本数、默认状态与边界一致；原文件保护和全页版面核验'],['下一研究因素','分析新图成功损伤及剩余行动机会，选一个可检验假设','先冻结同预算对照和停止条件，再执行；无明确SR收益则收束'],['扩展条件','选实际区域扩大或更真实目标条件中的一项','使用新的预留源文件；目标图贡献与SR/SG共同检查'],['可选演示','制作轨迹回放以展示观察、线索和动作','只读已保存轨迹；界面不改变环境、答案权限或评测规则'],['最终交付','课程总结、证据链、可运行入口与复现实验说明','区分仿真结果与实飞目标；报告负结果和资源预算']],[1.9,5.3,6.7])
        p(c,'新候选按对应冻结协议判断：最近探索升级采用SR至少增加2点、至少2/3正权重、SG不差；目标干预要求真实目标对控制至少增加5点。独立确认另要求源文件区间支持收益，并保护原协议性能。具体门槛以运行前协议为准，不能把单个SR达到60%当所有阶段通过。')
        p(c,'当前原S2/S3/S4和SwissView同模态迁移验收有效。本轮继续训练升级已收束；后续只在用户选择新因素后另立实验批次。报告记录截至2026年10月2日的实际完成工作。')
    pages.append(next_steps)
    def sources(c):
        p(c,'实验依据与参考资料','sub')
        p(c,'实验表格取自下列已审计原目录中的汇总、验收和复核文件。完整路径、来源SHA及材料哈希保存在“编制依据.json”；当前实验入口见根目录README与“项目导航/实验索引.md”。','note')
        for code,label,path in source_body(e):p(c,f'{code}  {label}：DATA/processed_data/{path}','note')
        p(c,'参考论文','sub')
        p(c,'[1] Sarkar A, Sastry S, Pirinen A, et al. GOMAA-Geo: GOal Modality Agnostic Active Geo-localization. arXiv:2406.01917, 2024.','note')
        p(c,'[2] Mi L, Béchaz M, Chen Z, et al. GeoExplorer: Active Geo-localization with Curiosity-Driven Exploration. 项目所附论文资料，2025.','note')
        p(c,'[3] Sun Y, Zhang Y, Zhu P. DynCur-Geo: Dynamic Curiosity Reward Shaping for Multimodal Active Geo-Localization. 项目所附论文资料，2026.','note')
        p(c,'三篇PDF位于“论文/”；这里只说明借鉴机制，不使用不同论文协议的SR与本地结果直接相减宣称提升。','note')
    pages.append(sources)
    anchor=big._tbl
    for content in pages:anchor=append_page(d,anchor,prototype,content)
    tail=OxmlElement('w:p');pr=OxmlElement('w:pPr');spacing=OxmlElement('w:spacing');spacing.set(qn('w:line'),'20');spacing.set(qn('w:lineRule'),'exact');pr.append(spacing);tail.append(pr);anchor.addnext(tail)
    d.core_properties.title=TITLE+' 中期报告';d.core_properties.subject='综合工程设计中期检查';d.core_properties.author='陈德昕 李世显 陈龙辉';d.core_properties.last_modified_by=''
    d.save(REPORT)


def plain_doc(title):
    d=Document();s=d.sections[0];s.page_width=Cm(21);s.page_height=Cm(29.7);s.top_margin=s.bottom_margin=Cm(2.2);s.left_margin=s.right_margin=Cm(2.4)
    normal=d.styles['Normal'];normal.font.name='宋体';normal.font.size=Pt(11);normal.paragraph_format.line_spacing=1.35;normal.paragraph_format.space_after=Pt(6)
    for sn,size in [('Title',20),('Heading 1',14),('Heading 2',12)]:
        st=d.styles[sn];st.font.name='宋体';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0)
    d.add_paragraph(title,'Title');return d


def summary(e):
    d=plain_doc('主动视觉地理定位阶段提升总结');d.add_paragraph('2026年10月2日')
    d.add_paragraph('平台与单智能体主方法已经完成正式S2/S3验收，并通过10×10密度扩展和SwissView同模态迁移。当前默认保留冻结Small256探索器与EdgeTargetCue视觉线索头；继续训练候选目标图证据通过，但独立新图SR收益不足，未替换默认。')
    d.add_paragraph('从工程修复到目标视觉线索','Heading 1')
    d.add_paragraph('先修复环境与训练实现，统一SR/SG、计划题数、逐步日志和源文件划分；随后逐项比较云端排序、记忆和治理，以及本地奖励塑形、动作筛选、目标监督和模型容量。早期val100的动作筛选均值达到66%，目标遮蔽仍为65.67%，不能据此证明有效利用目标图。')
    d.add_paragraph('完整导航训练确立NoTarget探索器后，可信邻接与方向解耦曾全部弃权；加入细边缘连续性才获得可靠的局部邻接证据。开发140题×3的真实目标SR由72.62%升到88.57%，正式S2同题真实目标为89.33%，禁用cue为73.00%。各段对应不同队列或训练设置，不能连接为一条受控累计提升曲线。')
    d.add_paragraph('主线已经验收的结果','Heading 1')
    t=d.add_table(rows=1,cols=4);t.style='Light Shading Accent 1';t.autofit=False
    widths=[3.4,6.2,3.0,3.6]
    for col,width in zip(t.columns,widths):col.width=Cm(width)
    for c,x in zip(t.rows[0].cells,['验收','协议与样本','SR','SG']):c.text=x
    for code,label in [('S2','S2'),('S3','S3'),('S4','S4密度'),('Swiss','SwissView')]:
        x=e[code];r=t.add_row().cells
        for c,v in zip(r,[label,f"{x['k']}×{x['k']}；{x['tasks']}题×3权重",f"{x['sr_mean']:.2%}",f"{x['sg_mean']:.3f}"]):c.text=v
    for row in t.rows:
        for c,width in zip(row.cells,widths):c.width=Cm(width)
    d.add_paragraph('三权重是三份训练参数；源文件分别为4、10、20和20。10×10只提高同一原图内密度，SwissView只验证连续航拍目标的同模态迁移。')
    d.add_page_break();d.add_paragraph('最近继续训练的收益与停止原因','Heading 1')
    d.add_paragraph('同预算Continue5和Adapt10保持架构及视觉头不变。Continue5在已见10×10开发队列取得79.73%/SG0.823，真实目标对禁用、均值、错图分别领先20.33、27.53、22.53点，目标证据通过；Adapt10仅比Continue5增加0.27点且SG更差，协议适配未升级。')
    d.add_paragraph('固定20新SwissView源文件后，5×5原M0为86.73%/0.334，Continue5为87.40%/0.310；10×10原M0为72.07%/1.230，Continue5为72.67%/0.868。新10×10 SR只+0.60点，1/3正权重，95%差区间[−0.87,+2.27]点；恢复185与损伤176，净9/1500。SG收益真实保留，但不足以证明稳定SR优势。')
    d.add_paragraph('两批10500新记录加1500按SHA复用Full，共12000轨迹/181282动作重放通过。未增加训练或云端调用；新确认40—59号源文件已消费。三份Continue5权重保留，原默认不切换，也不在本批继续调参数追求过线。')
    d.add_paragraph('方法边界与后续','Heading 1')
    d.add_paragraph('当前为训练过的单导航智能体，探索器和视觉头属于内部模块。云端反思/双角色同预算小验证未支持扩大协作。旋转、可靠性校准、账本保护、邻域覆盖等候选均按原批次保留，不删去负结果。')
    d.add_paragraph('已完成主线可支撑课程阶段报告；任意角度、真实异时、跨视角、实际覆盖扩大和无人机实飞仍待验证。后续先完善材料，再依据成功损伤与预算机会选择一个新因素，使用新源文件确认；GUI可用于演示，不是评测前置。')
    d.add_paragraph('材料入口为根目录README、项目导航/实验索引.md和绘图/README.md。旧阶段总结v1保持原样，反映2026年9月30日当时状态；v2为当前阅读版。')
    d.save(SUMMARY)
    paragraphs=[p.text for p in d.paragraphs if p.text.strip()]
    rows=['| 验收 | 协议与每权重任务 | SR | SG |','|---|---|---:|---:|']
    for code,label in [('S2','S2'),('S3','S3'),('S4','S4密度'),('Swiss','SwissView')]:
        x=e[code];rows.append(f"| {label} | {x['k']}×{x['k']}；{x['tasks']}题×3权重 | {x['sr_mean']:.2%} | {x['sg_mean']:.3f} |")
    paragraphs.insert(paragraphs.index('主线已经验收的结果')+1,'\n'.join(rows))
    text='\n\n'.join(paragraphs)
    (OUT/'阶段提升总结_v2.md').write_text('# 阶段提升总结 v2\n\n'+text+'\n','utf8')


NAVY='20362D';GREEN='356B54';LIGHT='F5F7F3';INK='243329';GRAY='61726A';OCHRE='B57B30'


def color(s):return PRGB.from_string(s)
def txt(slide,x,y,w,h,text,size=20,col=INK,bold=False):
    shape=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=shape.text_frame;tf.word_wrap=True
    tf.margin_left=tf.margin_right=Inches(.015);tf.margin_top=tf.margin_bottom=Inches(.025)
    for i,line in enumerate(text.split('\n')):
        p0=tf.paragraphs[0] if i==0 else tf.add_paragraph();p0.text=line;p0.font.name='Microsoft YaHei';p0.font.size=PPt(size);p0.font.color.rgb=color(col);p0.font.bold=bold;p0.space_after=PPt(9)
    return shape


def rect(slide,x,y,w,h,fill,line=None,rounding=False):
    sh=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounding else MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h));sh.fill.solid();sh.fill.fore_color.rgb=color(fill)
    if line:sh.line.color.rgb=color(line);sh.line.width=PPt(1)
    else:sh.line.fill.background()
    return sh


def image_fit(slide,path,x,y,w,h):
    from PIL import Image
    with Image.open(path) as im:ratio=im.width/im.height
    ww=min(w,h*ratio);hh=ww/ratio
    return slide.shapes.add_picture(str(path),Inches(x+(w-ww)/2),Inches(y+(h-hh)/2),width=Inches(ww),height=Inches(hh))


def make_slide(prs,title,index,subtitle=''):
    sl=prs.slides.add_slide(prs.slide_layouts[6]);sl.background.fill.solid();sl.background.fill.fore_color.rgb=color(LIGHT)
    txt(sl,.62,.43,12,.72,title,32,bold=True)
    if subtitle:txt(sl,.64,1.22,12,.55,subtitle,15,GRAY)
    txt(sl,.64,7.04,11.8,.27,'主动视觉地理定位  ·  综合工程设计中期汇报  ·  2026.10.02',10,GRAY)
    txt(sl,12.28,7.03,.45,.28,f'{index:02d}',10,GRAY)
    return sl


def ppt_table(slide,x,y,w,h,rows,widths):
    sh=slide.shapes.add_table(len(rows),len(rows[0]),Inches(x),Inches(y),Inches(w),Inches(h));tb=sh.table
    for i,cw in enumerate(widths):tb.columns[i].width=Inches(cw)
    for ri,row in enumerate(rows):
        for ci,t in enumerate(row):
            c=tb.cell(ri,ci);c.text=str(t);c.margin_left=c.margin_right=Inches(.11);c.margin_top=c.margin_bottom=Inches(.08);c.vertical_anchor=MSO_ANCHOR.MIDDLE
            c.fill.solid();c.fill.fore_color.rgb=color(GREEN if ri==0 else ('FFFFFF' if ri%2 else 'EAF0E9'))
            for p0 in c.text_frame.paragraphs:p0.font.name='Microsoft YaHei';p0.font.size=PPt(18 if ri==0 else 19);p0.font.bold=ri==0;p0.font.color.rgb=color('FFFFFF' if ri==0 else INK)
    return sh


def deck(e):
    prs=Presentation();prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5);prs.core_properties.title=TITLE+' 中期汇报';prs.core_properties.author='陈德昕 李世显 陈龙辉'
    sl=prs.slides.add_slide(prs.slide_layouts[6]);sl.background.fill.solid();sl.background.fill.fore_color.rgb=color(NAVY)
    txt(sl,.72,.69,7.8,.5,'综合工程设计中期汇报',19,'D4DFD5')
    txt(sl,.72,1.70,7.65,2.15,'基于主动探索的\n视觉地理定位方法\n设计与实现',37,'FFFFFF',True)
    txt(sl,.75,4.29,7.2,.83,'平台与单智能体主线已验收\n网格密度与同模态迁移已有实证',22,'D4DFD5')
    txt(sl,.75,6.27,7.7,.63,'陈德昕  李世显  陈龙辉   |   指导教师 王笛\n2026 年 10 月 2 日',15,'D4DFD5')
    cover=ROOT/'DATA/raw_data/SwissView/data/SwissView100/swisstopo_20.jpg'
    image_fit(sl,cover,8.65,1.12,3.95,4.92)
    txt(sl,8.72,6.19,3.9,.45,'航拍切图仿真  ·  固定预算搜索',14,'D4DFD5')
    sl.notes_slide.notes_text_frame.text='约30秒。强调课题是主动探索：行动改变下一步观察。此处为真实SwissView航拍资料示意，实验结论限于切图仿真。不要把当前成绩表述为无人机实飞。'
    sl=make_slide(prs,'任务是连续观察与自主搜索',2,'给定目标图，在有限预算内通过移动获得新信息')
    for i,(a,b) in enumerate([('观察','当前航拍图\n给定目标图'),('判断','目标线索\n访问历史'),('行动','上右下左\n消耗一步'),('反馈','下一局部图\n剩余预算')]):
        x=.73+i*3.17;rect(sl,x,2.03,2.76,1.68,'E4ECE4',GREEN,True);txt(sl,x+.19,2.19,2.4,.43,a,25,GREEN,True);txt(sl,x+.19,2.81,2.4,.86,b,20)
        if i<3:
            chevron=sl.shapes.add_shape(MSO_SHAPE.CHEVRON,Inches(x+2.81),Inches(2.62),Inches(.22),Inches(.42));chevron.fill.solid();chevron.fill.fore_color.rgb=color(OCHRE);chevron.line.fill.background()
    txt(sl,.8,4.35,12,.69,'成功：预算内首次到达目标网格。SG：所有终局的距离均值，成功记 0。',21)
    txt(sl,.8,5.48,12,.93,'执行策略只读双图与公开位置、预算、访问历史。\n真实目标坐标、真实距离、未访问图块保留在环境与评测端。',20,GRAY)
    sl.notes_slide.notes_text_frame.text='约40秒。基础协议5×5/B10，扩展10×10/B20。越界停留原地并耗步，最后一步到达也成功。训练奖励可以用真值，推理输入不含真值。'
    sl=make_slide(prs,'参考框架落实为可运行平台',3,'训练小模型符合正式任务要求；论文借鉴范围与完整复现区分说明')
    ppt_table(sl,.72,1.99,11.9,3.65,[['参考','本项目落地','当前范围'],['GOMAA-Geo','航拍切图、预算搜索、历史序列决策','小型GRU替代大序列模型'],['GeoExplorer','预测误差与好奇心探索对照','未稳定优于PBRS'],['DynCur-Geo','动态门控与奖励塑形思想对照','按本地预算适配，保留负结果']],[2.45,5.2,4.25])
    txt(sl,.78,5.98,11.8,.68,'当前主方法采用 PBRS 与目标视觉线索；不宣称完整复现三篇论文。',21,GREEN,True)
    sl.notes_slide.notes_text_frame.text='约40秒。正式任务要求主动探索平台，允许训练，多Agent可选。GOMAA-Geo的零样本跨模态泛化不等于没有训练；本系统目前只验证航拍目标同模态。资料来源为本地三篇论文和任务对应说明。'
    sl=make_slide(prs,'单 Agent 把探索与可信目标线索结合',4,'冻结图像编码器 · Small256探索器 · EdgeTargetCue视觉头 · 行动门控')
    image_fit(sl,ASSETS/'报告_Agent闭环.png',.75,1.79,11.9,4.25)
    txt(sl,.82,6.34,11.7,.45,'探索器保持覆盖，视觉头在有可信邻接证据时接管；这两个模块属于同一导航智能体。',17,GRAY)
    sl.notes_slide.notes_text_frame.text='约45秒。NoTarget指探索器目标通道为训练均值，不是全系统忽略目标。视觉头比较目标与当前图的语义、局部关系及边缘；最高方向概率≥0.50且合法、未访问才接管。其余按探索器。系统有训练。'
    sl=make_slide(prs,'评测证据按源文件与协议组织',5,'固定任务、全部权重、逐动作日志、输入哈希与重放复核')
    for x,num,label in [(.82,'151','Masa 本地源文件'),(4.97,'100','SwissView100 源文件'),(9.1,'3','每方法冻结训练权重')]:
        txt(sl,x,2.02,3.0,1.04,num,54,GREEN,True);txt(sl,x,3.2,3.55,.61,label,22)
    ppt_table(sl,.79,4.19,11.7,1.39,[['数据划分','依据'],['Masa','137训练 / 4验证 / 10测试；探索器109拟合、28开发留出'],['SwissView','累计48文件跑过导航，52尚未跑；预训练覆盖未知']],[2.5,9.2])
    txt(sl,.82,6.0,11.7,.64,'三权重、重复路线、两种网格都不增加独立地图；区间按源文件整组计算。',20,GRAY)
    sl.notes_slide.notes_text_frame.text='约40秒。Masa本地规模与论文完整规模不同，不能把本地SR直接减论文SR。源文件独立不等于严格地理无重叠。回放检查逐步神经决策和环境终止，避免只有汇总数字。'
    sl=make_slide(prs,'正式主线已完成四项验收',6,'不同队列分别报告，不合并成一个总成功率')
    rows=[['验收','协议 / 每权重题数','SR','SG']]
    for code,label in [('S2','S2稳定'),('S3','S3独立源图'),('S4','S4密度扩展'),('Swiss','SwissView迁移')]:
        x=e[code];rows.append([label,f"{x['k']}×{x['k']} / B{x['budget']} / {x['tasks']}题",f"{x['sr_mean']:.2%}",f"{x['sg_mean']:.3f}"])
    ppt_table(sl,.77,1.96,11.82,3.56,rows,[3.25,4.47,2.05,2.05])
    txt(sl,.81,5.96,11.7,.67,'10×10增加同一原图内的网格密度；SwissView验证连续航拍目标的同模态迁移。',20,GRAY)
    sl.notes_slide.notes_text_frame.text='约45秒。四行均三权重平均。源文件数4/10/20/20，S4源文件已见，SwissView原正式迁移20新源文件。S2/S3/S4验收通过和后续新候选失败分别记录；后者不回写前者。'
    sl=make_slide(prs,'目标图收益通过干预验证',7,'Continue5固定三份最终权重；已见20源文件，10×10/B20')
    image_fit(sl,ASSETS/'汇报_目标证据简图.png',.77,1.84,9.30,4.70)
    txt(sl,10.38,2.07,2.13,.49,'真实目标',20,GREEN,True);txt(sl,10.38,2.7,2.10,.92,'79.73%',35,GREEN,True)
    txt(sl,10.38,3.79,2.05,2.15,'对三个控制\n均领先超过5点\n三权重全正\nSG同时改善',18)
    sl.notes_slide.notes_text_frame.text='约45秒。禁用cue59.40%、均值52.20%、错误图57.20%。均值和错图同时作用于目标global/local/edge通道，错误图环境仍按真目标评测。4500新控制+1500复用Full，全6000条审计。证明此队列目标图有用，不等于继续训练稳定优于原默认。'
    sl=make_slide(prs,'新源文件未支持继续训练升级',8,'固定SwissView40—59号20新源文件；两个协议共用同一队列')
    image_fit(sl,ASSETS/'报告_新源文件确认.png',.75,1.86,9.30,4.65)
    txt(sl,10.32,2.00,2.35,.50,'10×10 SR',20,GREEN,True);txt(sl,10.32,2.60,2.25,.95,'+0.60点',31,GREEN,True)
    txt(sl,10.32,3.67,2.17,2.35,'1/3权重正收益\n95%差区间含零\nSG 1.230→0.868\n原默认保持',18)
    sl.notes_slide.notes_text_frame.text='约50秒。原M0 72.07%，Continue5 72.67%，源文件95%SR差区间[-0.87,+2.27]点。恢复185、损伤176，净9/1500。5×5保护通过但10×10未满足+2点/2正权重/CI正。6000新记录重放、2500图块重建通过；不继续调参，也不追加条件控制。'
    sl=make_slide(prs,'其他因素保留完整正负结果',9,'等信息与预算的对照决定是否扩大；多个候选未进入默认')
    ppt_table(sl,.78,1.99,11.78,3.94,[['路线','观察','决策'],['旋转/可靠性','精度提高可能损失召回与真实线索','保留边界，未稳定升级'],['云端多角色','同预算反思/双角色未增加稳定SR','支线收束'],['零调用账本','完整队列SR+1.27点，未达+2点','低成本备选'],['协议适配','Adapt10仅比Continue5多0.27点，SG更差','不扩大网络、不升级默认']],[3.2,5.45,3.13])
    txt(sl,.82,6.24,11.7,.47,'负结果解释策略权衡，仍是可追溯研究成果。',21,GREEN,True)
    sl.notes_slide.notes_text_frame.text='约45秒。多Agent不是没做过，已有小实验和接口修复后的T/D/G同预算闭环；当前证据不足扩大。保护规则修复损伤但丢失恢复；不可通过叠加更多规则保证SR。所有旧判定与源码快照保留。'
    sl=make_slide(prs,'当前成果有明确适用范围',10,'工程与仿真证据已经建立；真实部署问题仍需分项验证')
    for i,(title,body) in enumerate([('已完成','主动交互平台\n目标图与局部视觉决策\nS2/S3/S4与同模态迁移'),('待验证','真实异时与任意角度\n跨视角或其他目标模态\n更大区域与无人机实飞'),('后续原则','先分析恢复与损伤\n单因素、冻结实验协议\n用新的源文件作确认')]):
        x=.82+i*4.15;rect(sl,x,2.02,3.63,3.71,'E8EEE5');txt(sl,x+.22,2.28,3.2,.51,title,25,GREEN,True);txt(sl,x+.22,3.15,3.15,2.13,body,21)
    txt(sl,.83,6.18,11.6,.56,'界面可服务于轨迹演示；是否需要训练、多角色或GUI由任务证据决定。',20,GRAY)
    sl.notes_slide.notes_text_frame.text='约40秒。当前单Agent已满足继续推进的基础条件。暂不建设大系统或重复训练；先收束材料再选择新假设。GUI不是成功率验收前置；若用于答辩，应只读已有日志。'
    sl=make_slide(prs,'交付材料与下一步安排',11,'统一事实口径，把重要入口、复现实验与阶段结论分开整理')
    for i,(num,title,body) in enumerate([('01','报告与汇报','教师模板中期报告、阶段总结、演讲提示'),('02','证据与索引','58个实验批次原目录入口、哈希与复核'),('03','维护与下一步','保留默认；分析失败后另立单因素实验')]):
        y=1.98+i*1.43;txt(sl,.85,y,.76,.64,num,32,OCHRE,True);txt(sl,1.97,y,9.8,.51,title,24,GREEN,True);txt(sl,1.98,y+.62,9.8,.63,body,21)
    sl.notes_slide.notes_text_frame.text='约35秒。代码、数据、模型和原始实验路径保留。清理只删除可再生缓存和核查后的重复副本，并留恢复清单。当前最新到科研图35，架构/边界图2张。后续最终报告按已验证范围表述。'
    sl=prs.slides.add_slide(prs.slide_layouts[6]);sl.background.fill.solid();sl.background.fill.fore_color.rgb=color(NAVY)
    txt(sl,.83,.94,11.6,.51,'阶段结论',22,'D4DFD5')
    txt(sl,.84,2.11,11.2,2.1,'单 Agent 主线已有正式证据\n新因素按收益和独立确认推进',36,'FFFFFF',True)
    txt(sl,.88,4.54,11.1,1.21,'原默认保持  ·  SR优先  ·  保留SG收益与负结果\n下一步完善课程交付并选择一个有失败证据支持的新因素',22,'D4DFD5')
    txt(sl,.89,6.64,10.0,.4,'陈德昕  李世显  陈龙辉   |   指导教师 王笛',16,'D4DFD5')
    sl.notes_slide.notes_text_frame.text='约20秒。收束：主线已验收，新增继续训练没有稳定SR优势，因此不替换默认。课题完成的是仿真主动搜索平台和受控证据，真实多模态与实飞属于后续。进入问答。'
    prs.save(SLIDES)
    (SLIDES.parent/'演讲提示与答辩问答.md').write_text('''# 中期汇报演讲提示与答辩问答

12页建议约8分钟；各页备注区包含讲解重点，展示时优先说明任务、当前方法、正式验收与最近停止原因。

## 常见问题

1. 这是Agent吗？是。它依据连续局部观测和历史自主选择动作，环境执行并返回下一观测；探索器和视觉头属于同一Agent内部模块。
2. 为什么训练小模型？正式要求允许训练；约80万参数策略适合现有硬件，参考框架也包含训练。当前系统不能称整体零训练。
3. GOMAA-Geo与GeoExplorer用到了吗？借鉴切图、预算搜索、序列决策与好奇心/奖励对照；采用自己的小模型平台，没有完整复现原论文。
4. 为什么主方案NoTarget还说利用目标图？NoTarget只指探索器目标通道被遮蔽，专门的目标视觉头读取给定目标与当前图并提供可信邻接线索。
5. S2现在通过了吗？原边缘主线已通过S2、S3、S4及SwissView同模态迁移。Continue5新候选未升级，不影响原验收。
6. SG改善为什么不升级？新图SR仅+0.60点、1/3权重正，区间包含零。SG是辅助定位误差，不能替代预算内到达目标的成功率。
7. 多Agent做到什么程度？已做同预算小验证和消息/接口诊断闭环；已测配置没有足够收益，未进入正式默认。
8. 10×10代表覆盖更大地区吗？当前只把同一原图切得更密，真实地理面积没有扩大。
9. 500题×3是1500独立地图吗？不是。权重和路线重复不增加源文件；最近双协议共用20新图，bootstrap按源文件整组计算。
10. 和论文相比更好吗？不能直接比较不同数据划分、目标模态、预算与环境协议的SR；本项目只报告同题对照及已验证范围。
11. 实际无人机能用了吗？当前是航拍切图仿真，尚未接入实飞、飞控、真实异时与跨视角目标。
12. 下一步是什么？先完善课程材料；依据恢复/损伤和预算机会选择一个因素，运行前冻结对照，再使用新的源文件确认。

## 重要材料

- 当前总入口：../../README.md。
- 全部实验索引：../../项目导航/实验索引.md。
- 科研图与图注：../../绘图/README.md。
- 本PPT内容来自已审计结果，完整来源与SHA见中期报告目录编制依据.json。
''','utf8')


def main():
    e,paths=evidence();figures(e);build_report(e);summary(e);deck(e)
    basis=dict(cutoff='2026-10-02',source_sha256=paths,current_default='EdgeTargetCue',default_changed=False,new_training_steps=0,cloud_calls=0,new_navigation_evaluation=0,
        current_mainline_passed=['S2','S3','S4-density','SwissView-same-modality'],Continue5_independent_upgrade=False,
        cover_metadata_source=PROPOSAL.relative_to(ROOT).as_posix(),cover_metadata_copied_without_invention=True,
        revised_outputs={x.relative_to(ROOT).as_posix():digest(x) for x in (REPORT,SUMMARY,SLIDES)},
        template_columns=['一、设计目标与任务','二、完成情况','三、存在的主要问题','四、下一步需要完成的工作'],
        original_docx_backup='中期报告相关/历史版本/2026-09-30_中期报告_修订前.docx',visual_review_pending=True)
    (OUT/'编制依据.json').write_text(json.dumps(basis,ensure_ascii=False,indent=2)+'\n','utf8')
    (OUT/'编制说明.md').write_text('''# 中期材料编制说明

2026-10-02修订。保留教师中期报告模板的封面、人员登记表、四个必填栏目、页眉页脚及A4页面，按实际工作扩展正文。成员与导师信息取自用户已填写的最新版选题报告，原选题报告未修改。原中期报告存入历史版本，并保留整理前完整恢复备份。

本稿采用当前正式验收与最新候选结论，替换旧“66%基线、S2未通过”的现在时表述。原S2/S3/S4/SwissView主线已通过；Continue5目标证据通过但新源图SR升级未通过，原默认保持。各队列分数单列，不宣称整体零训练、多Agent成功或无人机实飞。

编制依据.json记录来源、协议与哈希；阶段总结v2、12页中期汇报及演讲问答使用同一口径。Word/PPT经本机Office导出校样并逐页核验；校样图片放在排版审阅，主阅读入口在本目录README及项目导航。
''','utf8')
    new_json(RECORD/'材料构建收据.json',basis)
    print(dict(report=str(REPORT),summary=str(SUMMARY),slides=str(SLIDES),slides_count=12),flush=True)


if __name__=='__main__':main()
