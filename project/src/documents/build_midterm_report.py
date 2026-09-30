"""Fill the course midterm template from verified local experiment records."""
from pathlib import Path
from copy import deepcopy
from hashlib import sha256
import json
import re
import zipfile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT=Path(__file__).resolve().parents[3]
OUTDIR=ROOT/'中期报告相关'
WORK=OUTDIR/'排版审阅'
BASE=WORK/'中期报告_模板解析.docx'
TITLE='基于主动探索的视觉地理定位方法设计与实现'
OUT=OUTDIR/f'中期报告_{TITLE}.docx'
TRAIN=ROOT/'DATA/processed_data/Masa/训练结果'


def load(path):return json.loads(path.read_text(encoding='utf-8'))
def digest(path):return sha256(path.read_bytes()).hexdigest()


def font(run,size=12,bold=False):
    run.font.name='宋体';run.font.size=Pt(size);run.bold=bold
    run.font.color.rgb=RGBColor(0,0,0)
    rf=run._element.get_or_add_rPr().rFonts
    for slot in ('ascii','hAnsi','eastAsia','cs'):rf.set(qn('w:'+slot),'宋体')


def clean(cell):
    for child in list(cell._tc):
        if child.tag!=qn('w:tcPr'):cell._tc.remove(child)
    cell.add_paragraph()


def unsnap(para):
    flag=OxmlElement('w:snapToGrid');flag.set(qn('w:val'),'0')
    para._p.get_or_add_pPr().append(flag)


def p(cell,text,kind='body',newpage=False):
    para=cell.paragraphs[0] if len(cell.paragraphs)==1 and not cell.paragraphs[0].text and not cell.tables else cell.add_paragraph()
    unsnap(para)
    fmt=para.paragraph_format
    fmt.space_before=Pt(0);fmt.space_after=Pt(3)
    fmt.line_spacing=1.5;fmt.widow_control=True
    fmt.keep_with_next=kind in ('title','sub','caption')
    fmt.page_break_before=newpage
    if kind=='body':fmt.first_line_indent=Pt(24)
    elif kind in ('title','sub'):
        fmt.first_line_indent=Pt(0);fmt.space_before=Pt(6);fmt.space_after=Pt(5)
    elif kind=='caption':
        para.alignment=WD_ALIGN_PARAGRAPH.CENTER;fmt.line_spacing=1.1;fmt.space_after=Pt(4)
    elif kind=='note':
        fmt.line_spacing=1.2;fmt.space_after=Pt(5)
    font(para.add_run(text),10.5 if kind in ('caption','note') else 12,kind in ('title','sub','caption'))
    return para


def table(cell,rows,widths):
    t=cell.add_table(rows=len(rows),cols=len(widths));t.autofit=False;t.alignment=WD_TABLE_ALIGNMENT.CENTER
    tw=t._tbl.tblPr.find(qn('w:tblW'));tw.set(qn('w:type'),'dxa');tw.set(qn('w:w'),str(round(sum(widths)*567)))
    for i,w in enumerate(widths):t.columns[i].width=Cm(w)
    borders=OxmlElement('w:tblBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'000000');borders.append(el)
    t._tbl.tblPr.append(borders)
    for ri,row in enumerate(t.rows):
        trpr=row._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
        if ri==0:trpr.append(OxmlElement('w:tblHeader'))
        for ci,c in enumerate(row.cells):
            c.width=Cm(widths[ci]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            margin=OxmlElement('w:tcMar')
            for edge,n in (('top',55),('bottom',55),('left',65),('right',65)):
                el=OxmlElement('w:'+edge);el.set(qn('w:w'),str(n));el.set(qn('w:type'),'dxa');margin.append(el)
            c._tc.get_or_add_tcPr().append(margin)
            para=c.paragraphs[0];para.paragraph_format.line_spacing=1.15
            unsnap(para)
            para.paragraph_format.space_after=Pt(0);para.paragraph_format.space_before=Pt(0)
            if ri==0 or ci>0:para.alignment=WD_ALIGN_PARAGRAPH.CENTER
            font(para.add_run(str(rows[ri][ci])),10.5,ri==0)
    spacer=cell.paragraphs[-1];spacer.paragraph_format.space_after=Pt(0);spacer.paragraph_format.line_spacing=Pt(2)
    unsnap(spacer)
    font(spacer.add_run(''),2)
    return t


def replace_text(para,text,size=15):
    para.clear();font(para.add_run(text),size)


def build():
    curiosity=load(TRAIN/'好奇心配对验证_v3/配对汇总.json')
    boundary=load(TRAIN/'合法动作协作验证_v1/配对汇总.json')
    target=load(TRAIN/'目标图利用验证_v1/配对汇总.json')
    assert abs(boundary['arms']['Boundary']['sr_mean']-.66)<1e-10
    assert abs(target['arms']['TargetPairs']['正确目标']['sr_mean']-.58)<1e-10
    d=Document(BASE);big=d.tables[0]
    # Preserve the cover, the teacher's four sections, merged metadata table,
    # headers/footers and page geometry. Only fill/expand the body cells.
    for para in d.paragraphs:
        t=para.text.strip()
        if re.match(r'^题\s+目：',t):
            replace_text(para,'题    目：'+TITLE,14)
            para.paragraph_format.first_line_indent=Pt(0)
            ind=para._p.get_or_add_pPr().find(qn('w:ind'))
            if ind is not None:
                for attr in ('firstLineChars','leftChars'):
                    ind.attrib.pop(qn('w:'+attr),None)
        elif t.startswith('小组人员'):replace_text(para,'小组人员：【待填写】')
        elif t.startswith('指导老师'):replace_text(para,'指导老师：【待填写】')
        elif t.startswith('填写时间'):replace_text(para,'填写时间：2026 年 9 月 30 日')
    def fill(c,text):
        clean(c);para=c.paragraphs[0];para.alignment=WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.line_spacing=1.15;font(para.add_run(text),10.5)
    fill(big.rows[0].cells[-1],TITLE)
    for row in big.rows[2:5]:
        seen=set()
        for c in row.cells:
            if c._tc in seen:continue
            seen.add(c._tc)
            if not c.text.strip():fill(c,'【待填写】')
    seen=set()
    for c in big.rows[5].cells:
        if c._tc in seen:continue
        seen.add(c._tc)
        if not c.text.strip():fill(c,'【待填写】')
    for row in big.rows[6:]:
        row.height=None;row.height_rule=WD_ROW_HEIGHT_RULE.AUTO
        pr=row._tr.get_or_add_trPr()
        for el in list(pr):
            if el.tag==qn('w:cantSplit'):pr.remove(el)
        c=row.cells[0];c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP;clean(c)

    c=big.rows[6].cells[0]
    p(c,'一、设计目标与任务','title')
    p(c,'本项目面向仅能获得局部航拍观测的目标搜索任务，设计并实现主动视觉地理定位实验系统。给定目标图像和起始位置，智能体通过逐步观察与移动，在有限步数内到达目标所在网格。中期工作重点是建立可信的环境与评测流程，完成可运行原型及受控对照，识别影响定位效果的主要因素。')
    p(c,'现阶段采用5×5网格、上右下左四方向动作和10步移动预算，初始曼哈顿距离分为4—8五档。越界移动停留原地并消耗一步；预算内首次到达目标即成功，最后一步到达也计为成功。执行端只能读取目标图、当前局部图及位置、剩余预算、访问记录等公开状态，不能读取真实目标坐标、距离或未访问图块。')
    p(c,'具体任务包括：完成数据清单及划分校验；统一成功率、最终距离和工程完成率的统计；接入视觉语言模型并逐项验证记忆、循环治理和层级搜索；保存完整日志并进行轨迹重放；在证据支持后开展独立源图及扩展实验。')
    p(c,'原选题拟探索零训练多智能体方法。当前已实现单VLM与规则模块，并新增本地小策略网络训练分支作为对照和机制诊断。该分支需要训练，不能作为原“零训练多智能体”目标已经完成的证明；多角色协作仍是后续任务。')

    c=big.rows[7].cells[0]
    p(c,'二、完成情况','title',True)
    p(c,'（一）数据、交互环境与统一评测','sub')
    p(c,'已完成MASA本地数据处理和任务清单校验。当前train、val、test分别包含137、4、10张源图，固定任务分别为3425、100、250题。按源图隔离划分，并核对图块数量、尺寸、索引和文件哈希。val100实际包含86种不同的“区域—起点—目标”组合，重复轮次不视作新增独立地图。')
    p(c,'已实现局部观测网格环境，明确移动边界、预算扣减和终止规则；重编码图像以去除文件名和元数据线索。任务答案仅保留在训练奖励构造或评测端。现有120项回归测试全部通过，覆盖动作边界、最后一步成功、划分冒充、观测泄漏防护、序列训练及目标替换等情况。')
    p(c,'统一采用SR衡量预算内到达目标的比例，SG衡量全部任务最终曼哈顿距离均值，Q衡量正常完成的比例。正式放行时，SR按计划题数计算；中断题不能算成功，SG按最大网格距离作保守补入。另记录重访率、越界率、分距离结果和调用时延，避免用“程序完成”替代“定位成功”。')
    p(c,'已建立批次预登记、配置与源码快照、数据及模型哈希、逐步动作和终态日志。云端请求错误单列，不以随机动作掩盖失败；本地模型评测逐题重新执行神经决策及环境转移，核对日志和统计值，为后续单因素实验提供统一依据。')
    p(c,'（二）最小VLM接入与小规模策略实验','sub')
    p(c,'已通过云端多图接口接入gemma4:31b，每步输入目标图、当前图及公开状态。完成动作输出解析、请求日志和网络传输兼容修复，并在固定20道验证题上依次测试动作排序、循环治理、空间记忆、邻域确认和粗区域层级搜索。各因素独立保存配置，未将新增模块同时叠加后归因。')

    p(c,'表1  云端VLM路线的小规模对照','caption',True)
    table(c,[['方法','成功/计划题','SR','阶段观察'],
        ['直接输出动作（A）','1/20','5%','往返循环较多'],
        ['四动作排序（B）','2/20','10%','收益有限'],
        ['排序＋循环治理（C）','2/20','10%','重访减少，SR未提升'],
        ['空间记忆（E）','3/20','15%','正常完成19题'],
        ['目标邻域确认（F）','1/20','5%','负结果，冻结'],
        ['粗区域层级搜索（G）','10/20','50%','与同题Frontier持平'],
        ['无图Frontier规则','10/20','50%','覆盖效率较高']], [5.1,2.25,1.25,5.3])
    p(c,'注：表中均按计划20题计算SR，E、F各有1题API中断。最初一次同设置最小VLM试验为0/20；A为后续独立批次。小规模结果不与val100直接排名。','note')
    p(c,'G采用3×3粗区域排序后再选择四方向动作，20题SR为50%、SG为2.00，但尚未显示相对无图规则的SR优势。完整val100的云端S2验证遇到超时、连接异常和限流，G、E首轮分别仅20/100题正常完成。批次收尾不等于评测完成，因此不能据此宣布S2通过。')
    p(c,'（三）本地小策略网络与训练实现修正','sub')
    p(c,'本地分支采用冻结Sat2Cap特征与约80万参数的GRU策略，输入目标、当前图块特征及公开状态，输出四方向动作。在已有CUDA环境中使用PPO与势函数奖励塑形（PBRS）训练；目标距离仅用于构造奖励，不作为执行输入。')
    p(c,'已修正早期训练的到达奖励、终态势函数和序列梯度问题。旧试跑值仅作历史记录；后续统一初始化、训练步数和最终权重，采用三个续训种子配对。')

    p(c,'表2  本地val100上的主要对照结果','caption',True)
    rows=[['方法','平均SR','平均SG','处理意见']]
    def add(label,m,decision):rows.append([label,f"{m['sr_mean']:.2%}",f"{m['sg_mean']:.3f}",decision])
    add('修正PBRS',curiosity['arms']['PBRS'],'保留训练对照')
    add('PBRS＋好奇心',curiosity['arms']['Curiosity'],'未见收益')
    add('PBRS＋好奇心＋距离门控',curiosity['arms']['GatedCuriosity'],'未超过PBRS')
    add('PBRS＋合法动作筛选',boundary['arms']['Boundary'],'当前保留方案')
    add('再加目标方向监督',target['arms']['TargetPairs']['正确目标'],'未通过候选验收')
    add('Frontier无图规则',boundary['rules']['Frontier'],'最强已测规则')
    add('固定区域顺序规则',boundary['rules']['FixedRegion'],'规则对照')
    table(c,rows,[5.8,2.0,2.0,4.1])
    p(c,'注：学习方法为三个续训种子的平均值，各次训练40,960个实际环境步；确定性规则执行一次并核验可复现。三种子共享历史预训练初始化，不是三次独立预训练。所有方法使用同一5×5、B=10、C=4—8的val100清单。','note')
    p(c,'好奇心实验采用独立预训练、动作条件化且在PPO阶段冻结的预测器，误差在train内归一化。距离门控虽改善无门控好奇心，但未超过PBRS，故未继续叠加；该适配不等于完整复现DynCur-Geo。')
    p(c,'合法动作筛选只读取公开行列位置，排除向网格外移动的方向，并在训练采样、PPO概率更新和推理时保持一致。该因素将平均SR由44.33%提高到66.00%，SG由1.507降至1.020，三个种子均提升，越界率降到0；相对同题Frontier高13个百分点。环境本身的越界耗步规则未改变。')
    p(c,'三个种子为66%、72%、60%，采用66%的均值而非最高成绩。合法动作筛选是已有组件，该结果作为工程改进和后续基线，不独立宣称算法创新。')

    p(c,'（四）目标信息利用验证与阶段产出','sub',True)
    p(c,'已分别输入真实目标、固定train特征均值和错误目标图，独立执行完整轨迹，错误图仍按原真实目标计分。动作筛选方案在遮蔽目标后仍有65.67%的SR，仅比真实目标低0.33个百分点，目标贡献证据不足。')
    p(c,'进一步加入成对目标方向监督：相同3帧观察历史对应两个不同目标，模型学习各自正确方向。模型结构、PPO奖励和训练步数不变，辅助损失权重固定0.5；目标位置仅用于train端生成标签。全部训练结束后再评测val，不扫描参数。')
    p(c,'表3  目标图干预结果（val100，三个种子平均）','caption')
    table(c,[['方法','真实目标SR','遮蔽目标SR','错误目标SR'],
        ['保留的动作筛选方案','66.00%','65.67%','62.00%'],
        ['加入目标方向监督','58.00%','61.00%','63.00%']], [5.5,2.8,2.8,2.8])
    p(c,'新监督没有通过候选验收：平均SR下降8个百分点，且真实目标未优于遮蔽和错误目标。成对决策诊断中“两目标动作均正确”率在train源图上由17.90%升至31.97%，在val源图上仅由14.97%升至17.32%，未转化为完整定位收益。因此保留原动作筛选方案，将新监督配置归档为负结果。')
    p(c,'已交付环境、VLM接口、规则策略、训练与消融脚本及对应模型、快照和报告。好奇心配对、动作筛选、目标监督三批分别保存1800、2300、1800条评测轨迹，均通过批内重放审计。各批共享任务，这5900条轨迹并非5900个独立地理任务。')
    p(c,'平台建设、原型实现和问题定位已完成；多VLM协作、完整S2、独立测试及大网格与跨数据集实验尚未完成。')

    c=big.rows[8].cells[0]
    p(c,'三、存在的主要问题','title',True)
    p(c,'1. 目标视觉信息尚未形成稳定正贡献。动作筛选方案在真实和遮蔽目标下的SR接近，增加方向监督也未改善定位。这说明目前还不能把66%的结果归因于有效的目标识别与视觉引导。输入干预可能带来分布变化，因此也不能直接断言模型完全没有使用视觉。后续应先检验特征匹配的跨源图泛化，再接回导航。')
    p(c,'2. 结果仍不满足完整阶段验收。S2要求平均SR≥60%、SG≤1.8，并同时约束相对规则收益、分距离表现、每轮重访率、三轮波动和组件消融。动作筛选方案虽然达到SR、SG要求，但三轮极差为12个百分点，超过10点；一轮重访率为11.38%，超过10%；目标贡献未达预定的5点增益，故尚不能进入扩展阶段。')
    p(c,'3. 数据规模与统计独立性有限。val仅4张源图，100题包含重复路线。继续在同一验证集上调整容易高估泛化能力。与论文相比，本地划分、训练池、预算和边界处理也不完全相同；论文中的成功率只能参考量级，不能直接相减后宣称优于论文。')
    p(c,'4. 原选题目标与已实现内容仍有差距。当前系统主要是单VLM或小模型与规则模块配合，尚不具备经过验证的规划、执行和检查多角色协作。新增本地训练分支有助于受控诊断，但不能称为零训练方案。后续需围绕原课题明确协作模块的必要性、信息边界及调用成本。')
    p(c,'5. 云端验证和旧实现影响了实验进度。完整云端S2遇到超时、连接异常和限流；早期训练错误使部分结果只能作为历史记录。虽然已修正实现并补齐审计，后续仍需坚持预登记、统一预算和等条件对照，避免同时增加多个因素、重复试到过线或选最好种子报告。')

    p(c,'四、下一步需要完成的工作','title',True)
    p(c,'后续按“先验证目标信息，再处理搜索行为，最后扩展”的顺序推进。当前保留PBRS＋合法动作筛选作为主基线，冻结本轮未通过的目标方向监督配置，不将其继续叠加到新方案中。')
    p(c,'表4  后续工作与检查点','caption')
    table(c,[['顺序','拟完成工作','检查点与产出'],
        ['第一阶段','train内部按源图留出，检验目标与观察的匹配泛化；比较真实、遮蔽及替换目标。','固定划分与预算，先完成匹配诊断，再决定是否接回导航。'],
        ['第二阶段','目标信息有证据后，分别处理重复访问和训练波动。','同题、同预算、三种子配对；分别报告效果与行为指标。'],
        ['第三阶段','实现最小规划/执行或检查角色协作，并做在线消融。','只共享公开观测；增加等调用成本单Agent对照，统计请求与时延。'],
        ['第四阶段','冻结方法与正式S2协议，通过后做S3独立源图复核。','完成val100×3及必要消融，达标后使用test250；不据test反复调参。'],
        ['第五阶段','S3通过后，逐项扩展10×10及跨数据集实验，完成总结。','分别登记预算、清单和基线，保留负结果并交付代码与报告。']], [1.8,6.0,6.1])
    p(c,'候选因素保留要求平均SR相对同轮对照至少提高5个百分点、SG不变差，并在至少2/3种子中为正增益。涉及目标信息的方案还应在真实目标相对遮蔽及错误目标的在线评测中体现稳定收益，不能只依赖动作敏感性或训练集准确率。')
    p(c,'未达检查点时继续分析和修复，不把实验完成写成方法通过；最终报告呈现问题、改动、对照证据与适用边界。')
    p(c,'本报告依据截至2026年9月30日的项目日志、配对汇总及审计记录编制。','note')

    # Word ignores paragraph pageBreakBefore inside a spanning table row.
    # Split only the expanded full-width content into continuation tables,
    # keeping teacher borders/grid and all four section labels unchanged.
    chunks=[]
    starts=('二、完成情况','表1  ','表2  ','（四）目标信息','三、存在','四、下一步')
    for row in big.rows[7:]:
        for child in list(row.cells[0]._tc):
            if child.tag==qn('w:tcPr'):continue
            text=''.join(child.xpath('.//w:t/text()'))
            if text.startswith(starts):chunks.append([])
            if chunks:chunks[-1].append(deepcopy(child))
    prototype=deepcopy(big.rows[7]._tr)
    for row in list(big.rows)[7:]:big._tbl.remove(row._tr)
    anchor=big._tbl
    for chunk in chunks:
        page=d.add_paragraph();page.paragraph_format.space_after=Pt(0);page.paragraph_format.space_before=Pt(0)
        unsnap(page)
        page.paragraph_format.line_spacing=Pt(1)
        page.paragraph_format.page_break_before=True
        page.paragraph_format.keep_with_next=True
        font(page.add_run(''),1)
        anchor.addnext(page._p)
        new=OxmlElement('w:tbl');new.append(deepcopy(big._tbl.tblPr));new.append(deepcopy(big._tbl.tblGrid))
        tr=deepcopy(prototype);tc=tr.find(qn('w:tc'))
        for ch in list(tc):
            if ch.tag!=qn('w:tcPr'):tc.remove(ch)
        for ch in chunk:
            for flag in ch.xpath('.//w:pageBreakBefore'):flag.getparent().remove(flag)
            tc.append(ch)
        if tc[-1].tag!=qn('w:p'):tc.append(OxmlElement('w:p'))
        new.append(tr);page._p.addnext(new);anchor=new

    # The mandatory final paragraph must not force a blank final page.
    tail=d.paragraphs[-1]
    if not tail.text:
        tail.paragraph_format.line_spacing=Pt(1)
        tail.paragraph_format.space_before=Pt(0);tail.paragraph_format.space_after=Pt(0)
        unsnap(tail);font(tail.add_run(''),1)

    # Explicit font slots for all inherited cover/metadata/header/footer runs.
    for part in [d.part]+[s.header.part for s in d.sections]+[s.footer.part for s in d.sections]+[s.first_page_header.part for s in d.sections]+[s.first_page_footer.part for s in d.sections]:
        for r in part.element.xpath('.//w:r'):
            rp=r.get_or_add_rPr();rf=rp.rFonts
            if rf is None:rf=OxmlElement('w:rFonts');rp.insert(0,rf)
            for slot in ('ascii','hAnsi','eastAsia','cs'):rf.set(qn('w:'+slot),'宋体')
    d.core_properties.title=TITLE+'：中期报告'
    d.core_properties.subject='综合工程设计中期检查'
    d.core_properties.author='';d.core_properties.last_modified_by=''
    OUTDIR.mkdir(exist_ok=True);d.save(OUT)
    evidence=[ROOT/'课程相关报告模版/综合工程设计-2-中期报告.doc',
        ROOT/'README.md',ROOT/'选题报告相关/分阶段验收标准_v1.md',
        TRAIN/'好奇心配对验证_v3/配对汇总.json',TRAIN/'合法动作协作验证_v1/配对汇总.json',
        TRAIN/'目标图利用验证_v1/配对汇总.json',TRAIN/'目标图利用验证_v1/独立复核.json',
        ROOT/'DATA/processed_data/Masa/评测结果/S2稳定验证_v1/S2验证报告.md']
    ledger={str(x.relative_to(ROOT)):digest(x) for x in evidence}
    (OUTDIR/'编制依据.json').write_text(json.dumps(dict(cutoff='2026-09-30',evidence_mode='existing real local experiment logs; no invented results',
        source_sha256=ledger,word_sha256=digest(OUT),template_columns=['一、设计目标与任务','二、完成情况','三、存在的主要问题','四、下一步需要完成的工作'],
        missing=['小组成员姓名','学号','联系方式','指导教师姓名及联系方式']),ensure_ascii=False,indent=2),encoding='utf-8')
    (OUTDIR/'编制说明.md').write_text('''# 中期报告编制说明

报告严格沿用教师中期模板的封面、人员登记表、四个栏目、页眉页脚及A4页面设置，正文按实际进展扩展。原始.doc模板未修改；Word成稿位于本文件同目录，排版审阅目录内的PDF及图片仅用于校样。

正文依据已有真实运行结果编制，未生成模拟数据、未重新训练或追加模型调用。相对于选题报告，本稿改为已实现系统、受控实验、负结果与后续检查点，不将原计划写成已完成。66%为三个续训种子均值，不是独立测试集成绩；目标信息贡献和完整S2尚未通过。

封面和登记表中的小组成员、学号、联系方式及指导教师信息在现有材料中均为空，保留【待填写】。请提交前补齐这些个人字段，实验内容无需替换。模板未要求实验截图，报告采用可追溯结果表和分析说明，不插入占位截图。
''',encoding='utf-8')
    print(json.dumps(dict(output=str(OUT),bytes=OUT.stat().st_size),ensure_ascii=False))


if __name__=='__main__':build()
