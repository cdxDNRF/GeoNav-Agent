"""New scientific figures and a new midterm summary; never overwrites user drafts.

Figures: existing research Python with matplotlib. DOCX: bundled artifact Python.
"""
import argparse
import hashlib
import json
from pathlib import Path

ROOT=Path(__file__).resolve().parents[3]
FIG=ROOT/'绘图/平台方法图_v1'
REPORT=ROOT/'总结报告相关/平台交付_v1'
RESULT=ROOT/'DATA/processed_data/MassGIS/评测结果/平台单模型对照_v1/主对照/同题批量结果_v1.json'


def sha(path):return hashlib.sha256(path.read_bytes()).hexdigest()
def read(path):return json.loads(path.read_text(encoding='utf-8-sig'))


def configure_plot():
    import matplotlib
    matplotlib.use('Agg')
    from matplotlib import font_manager
    font=Path('C:/Windows/Fonts/msyh.ttc')
    if font.exists():
        font_manager.fontManager.addfont(str(font))
        matplotlib.rcParams['font.family']=font_manager.FontProperties(fname=str(font)).get_name()
    matplotlib.rcParams.update({'pdf.fonttype':42,'ps.fonttype':42,'svg.fonttype':'none','font.size':11,'axes.unicode_minus':False})


def export(fig,name):
    for extension in ('png','svg','pdf'):
        path=FIG/f'{name}.{extension}'
        if path.exists():raise ValueError('图原件已存在，不覆盖：'+str(path))
        fig.savefig(path,dpi=300,bbox_inches='tight',facecolor='white')


def method_figures():
    configure_plot()
    import matplotlib.pyplot as plt
    from matplotlib.patches import FancyBboxPatch
    blue='#2467A4';green='#268773';orange='#B77720';gray='#566573'
    def canvas(title):
        f,a=plt.subplots(figsize=(13,7));a.set_xlim(0,13);a.set_ylim(0,7);a.axis('off')
        a.text(.2,6.65,title,fontsize=19,weight='bold');return f,a
    def box(a,x,y,w,h,text,color):
        a.add_patch(FancyBboxPatch((x,y),w,h,boxstyle='round,pad=.08,rounding_size=.1',facecolor=color+'18',edgecolor=color,lw=1.4))
        a.text(x+w/2,y+h/2,text,ha='center',va='center',fontsize=11,color='#162C3A',linespacing=1.6)
    def arrow(a,x,y,xx,yy,text=''):
        a.annotate('',xy=(xx,yy),xytext=(x,y),arrowprops=dict(arrowstyle='->',color=gray,lw=1.3))
        if text:a.text((x+xx)/2,(y+yy)/2+.12,text,ha='center',fontsize=9,color=gray)
    f,a=canvas('主动探索视觉地理定位  训练与实时导航闭环')
    a.text(.2,6.05,'离线训练与冻结',fontsize=12,color=gray)
    box(a,.3,4.9,2.8,.88,'训练 / 校准源\n按源隔离',orange)
    box(a,3.55,5.4,5.65,.45,'PBRS奖励塑形 + PPO：Small256 GRU探索器',orange)
    box(a,3.55,4.65,5.65,.45,'独立视觉头训练：语义 / 边缘连续性、邻接 / 负样本',orange)
    box(a,9.75,4.9,2.8,.88,'权重 / 均值冻结\n输入与源码SHA绑定',orange)
    arrow(a,3.1,5.34,3.45,5.62);arrow(a,3.1,5.15,3.45,4.86)
    arrow(a,9.3,5.62,9.65,5.34);arrow(a,9.3,4.86,9.65,5.15)
    a.text(.2,4.4,'在线推理与环境交互',fontsize=12,color=gray)
    box(a,.3,3.05,2.25,.95,'当前航拍观察\n公开位置 / 历史 / 预算',blue)
    box(a,3.15,3.05,2.35,.95,'Sat2Cap表征\n冻结编码缓存',blue)
    box(a,6.1,3.05,2.35,.95,'GRU实时推理\n探索器目标通道均值遮蔽',orange)
    box(a,9.1,2.85,3.2,1.15,'合法 / 新格门控\n目标线索优先\n覆盖收益与径向保护',green)
    arrow(a,2.6,3.52,3.05,3.52);arrow(a,5.55,3.52,6,3.52);arrow(a,8.5,3.52,9,3.52)
    box(a,.3,1.45,2.25,.9,'给定航拍目标图\n始终由视觉头使用',blue)
    box(a,3.15,1.45,5.3,.9,'当前 / 目标语义特征 + 20维边缘连续性\n目标视觉头实时推理  →  五类概率',orange)
    arrow(a,2.6,1.9,3.05,1.9);arrow(a,8.5,1.9,10.7,2.75)
    arrow(a,4.3,3,4.3,2.45)
    box(a,9.1,1.45,3.2,.9,'上 / 右 / 下 / 左\n环境执行并返回新观察',green)
    arrow(a,10.7,2.75,10.7,2.45)
    a.plot([10.7,10.7,.15,.15],[1.36,.72,.72,3.52],color=gray,lw=1.2)
    arrow(a,.15,3.52,.28,3.52)
    a.text(4,.9,'反馈闭环：新观察 / 公开历史；跨episode清空记忆',fontsize=9,color=gray)
    a.text(.3,.2,'策略不读取目标真值、全图或未访问图块；目标坐标只由环境用于终局评测。\n本地系统有训练来源；当前是单导航Agent，规则模块不等于多个自主Agent。',fontsize=10,color=gray)
    export(f,'01_训练与实时导航闭环');plt.close(f)
    f,a=canvas('决策机制  目标线索优先与覆盖保护')
    box(a,.35,4.65,5.4,1.18,'目标线索门控\n五类softmax最高方向 / 阈值0.50 / 合法且未访问\n通过则执行目标线索动作',blue)
    box(a,6.35,4.65,5.9,1.18,'无已接受目标线索\n最多三步合法路径：新增可行动覆盖机会\n相对原提议收益至少增加2',green)
    box(a,.35,2.5,5.4,1.25,'径向保护\nρ(v) = |r(v)−r(start)| + |c(v)−c(start)|\n候选下一格的ρ不小于原动作下一格',green)
    box(a,6.35,2.5,5.9,1.25,'动作优先级\n已接受cue → 受保护的覆盖候选 → 原GRU提议\n不能读真实目标距离，不新增stop动作',orange)
    arrow(a,8.3,4.5,3,3.9);arrow(a,5.85,3.1,6.2,3.1)
    a.text(.5,1.65,r'$\Delta G=G(p_{candidate})-G(p_{original})\geq 2$',fontsize=19,color=green)
    a.text(.5,.9,'协议参数：GRU hidden=256 · cue阈值=0.50 · 覆盖深度≤3 · 最小收益=2\n已确认区域默认：MassachusettsRoads，10×10 / B20 / 300米每格\n本平台：MassGIS已消费开发区5×5/B10与10×10/B20；5格只用M0，无覆盖治理。',fontsize=11,color=gray)
    export(f,'02_决策公式与参数');plt.close(f)


def result_figure():
    configure_plot()
    import matplotlib.pyplot as plt
    data=read(RESULT);names=list(data['metrics']);values=[data['metrics'][n] for n in names]
    f,axes=plt.subplots(1,2,figsize=(11,4.5));colors=['#2467A4','#D89242']
    for ax,key,title in zip(axes,['sr_planned','sg_m'],['预算内成功率 SR','全部终局距离 SG']):
        metrics=[v[key]*100 if key=='sr_planned' else v[key] for v in values]
        if any(v is None for v in metrics):
            ax.text(.5,.5,'队列未完整完成\n不计算完整SG',ha='center',va='center',transform=ax.transAxes)
        else:
            ax.bar(names,metrics,color=colors,width=.5)
            for x,v in enumerate(metrics):ax.text(x,v+(.8 if key=='sr_planned' else 8),f'{v:.1f}',ha='center',fontsize=13)
            ax.set_ylim(0,100 if key=='sr_planned' else max(metrics)*1.23+10)
        ax.set_title(title);ax.spines[['top','right']].set_visible(False);ax.grid(axis='y',alpha=.2);ax.set_axisbelow(True)
        ax.set_ylabel('%' if key=='sr_planned' else '米（成功计0）')
    cloud=data['metrics'].get('Cloud',{})
    f.suptitle('已消费开发区20题同题对照  5×5 / B10 / seed0',fontsize=16)
    f.text(.5,.015,f"单地区、不同策略比较；非独立确认、无CI。Cloud请求{cloud.get('api_requests',0)}，回退{cloud.get('fallback_actions',0)}；不据此升级默认。",ha='center',fontsize=10)
    f.tight_layout(rect=(0,.06,1,.93));export(f,'03_单模型同题工程对照');plt.close(f)


def source_manifest():
    selected=[ROOT/'project/local_policy_default.json',ROOT/'project/defaults/masa_roads_grid10_radial_v1.json',
        ROOT/'project/src/agents/parameterized_coverage.py',ROOT/'project/src/agents/massgis_navigator_v1.py',Path(__file__)]
    if RESULT.exists():selected.append(RESULT)
    value=dict(inputs_sha256={p.relative_to(ROOT).as_posix():sha(p) for p in selected},
        figures_sha256={p.relative_to(ROOT).as_posix():sha(p) for p in FIG.iterdir() if p.suffix in ('.png','.pdf','.svg')},
        current_system='trained single navigation agent',independent_geographic_units_in_new_comparison=1,
        default_changed=False,ground_or_text_goal_supported=False)
    path=FIG/'来源/图表输入与输出绑定_v1.json'
    if path.exists():raise ValueError('绑定已存在')
    path.write_text(json.dumps(value,ensure_ascii=False,indent=2),encoding='utf-8')


def document():
    from docx import Document
    from docx.shared import Cm,Pt
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    result=read(RESULT);d=Document();sec=d.sections[0]
    sec.page_width=Cm(21);sec.page_height=Cm(29.7);sec.top_margin=Cm(2.2);sec.bottom_margin=Cm(2.1);sec.left_margin=Cm(2.2);sec.right_margin=Cm(2.2)
    normal=d.styles['Normal'];normal.font.name='Microsoft YaHei';normal.font.size=Pt(10.5)
    normal._element.rPr.rFonts.set(qn('w:eastAsia'),'微软雅黑')
    normal.paragraph_format.space_after=Pt(7);normal.paragraph_format.line_spacing=1.2
    for name in ('Title','Heading 1','Heading 2'):
        style=d.styles[name];style.font.name='Microsoft YaHei';style.font.color.rgb=__import__('docx').shared.RGBColor(0,0,0)
        style._element.rPr.rFonts.set(qn('w:eastAsia'),'微软雅黑')
    d.add_heading('主动探索视觉地理定位中期平台交付总结',0)
    d.add_paragraph('2026年10月5日  平台阶段交付  基于当前冻结代码与实验记录')
    d.add_paragraph('本总结说明当前方法、实时平台、单云端模型工程对照及适用边界。项目已有训练过的单导航Agent和正式分阶段证据；本轮将历史回放扩展为真实策略运行、可配置API、同题批量和方法展示。多Agent协作、新训练及备用区域验证均未纳入本轮。中期平台交付不表示所有跨数据集或大范围研究均通过，课程最终验收仍以教师要求为准。')
    d.add_heading('研究进度与交付范围',1)
    table=d.add_table(rows=1,cols=4);table.style='Light Shading Accent 1'
    for cell,text in zip(table.rows[0].cells,['阶段','协议','SR','结论']):cell.text=text
    for row in [('原S2','5×5/B10','89.33%','原协议通过'),('原S3','5×5/B10','86.80%','原源图确认通过'),('原S4','10×10/B20密切图','77.13%','原密度扩展通过'),('SwissView','5×5/B10','86.80%','同模态迁移通过'),('Roads真实9km²','300米/格、B20','68.04%','径向保护确认通过'),('MassGIS正式','10×10 / 15×15','54.62% / 20.27%','新域必要门槛未过')]:
        for cell,text in zip(table.add_row().cells,row):cell.text=text
    d.add_paragraph('这些成绩来自不同队列，不能连接成累计提升曲线或直接宣称优于论文。前三权重平均不等于seed0单次成绩；Roads9km²默认seed0在原确认批次为64.80%。')
    d.add_heading('本轮完成内容',1)
    for text in ['本地实时：已消费开发区5×5/B10与10×10/B20，选择任务/策略/权重后初始化、单步、自动、暂停和结束；未访问图块遮挡，终止后导出原件。','单模型API：兼容chat/completions，支持可选Key；后端内存保存，双图和公开状态输入，结构检查、失败回退、请求上限与熔断。Gemma沿用既有云端配置。','批量与展示：固定同题队列，按计划分母计算SR，SG含成功0；方法说明与源码/输入绑定，保留成功与失败和完整运行日志。']:
        d.add_paragraph(text,style='List Bullet')
    d.add_page_break();d.add_heading('方法与训练推理结构',1)
    d.add_paragraph('冻结Sat2Cap提供图像表征。Small256 GRU探索器由PBRS奖励塑形与PPO训练；探索器目标通道以训练均值遮蔽，但目标视觉头继续接收给定目标与当前观察的语义/局部及20维边缘连续性特征。系统整体有训练来源，不是零训练，也不是多个自主Agent。')
    d.add_picture(str(FIG/'01_训练与实时导航闭环.png'),width=Cm(16.4))
    d.add_paragraph('图1  训练与在线推理分开。平台复用封存图像编码缓存，而GRU和视觉头每一步真正forward；历史回放模式不推理。目标真值只供环境终局评测，不进入策略。')
    d.add_heading('决策依据',1)
    d.add_paragraph('先检查目标头的最高方向是否满足0.50阈值、合法且未访问；若接受，cue优先。否则在最多三步合法路径中选择增加覆盖机会的候选，收益至少比原提议多2，且下一步到公开起点的曼哈顿径向距离不小于原动作。保护未过时保留原GRU提议。该规则不读取真实目标距离。')
    d.add_page_break();d.add_heading('参数公式与适用协议',1)
    d.add_picture(str(FIG/'02_决策公式与参数.png'),width=Cm(16.4))
    d.add_paragraph('图2  ρ为相对公开起点的距离，ΔG为候选与原提议的覆盖机会收益差。最高cue被拒绝时不重新排名其它cue；基础动作仅上、右、下、左，越界原地且耗一步，最后一步到达仍判成功。')
    table=d.add_table(rows=1,cols=2);table.style='Light Shading Accent 1'
    for cell,text in zip(table.rows[0].cells,['参数','值或含义']):cell.text=text
    for name,value in [('GRU隐藏维度','256'),('cue接受阈值','0.50（原冻结）'),('覆盖规划深度','min(3,剩余预算)'),('最小覆盖收益','2'),('径向参考','episode公开起点'),('当前已确认默认','Roads10×10/300米/B20'),('本轮开发界面','MassGIS img_6100；5格M0，10格M0/Coverage'),('API初版','Gemma；30秒、512输出tokens、temperature0')]:
        for cell,text in zip(table.add_row().cells,[name,value]):cell.text=text
    d.add_paragraph('5格没有Coverage治理；MassGIS使用现有冻结权重做已消费区域工程演示，不意味着其正式新域失败结论被改变。')
    d.add_page_break();d.add_heading('单模型同题工程对照',1)
    d.add_paragraph('固定DATA007原5×5任务前20题、seed0、B10，仅一个已消费地理区域。本地M0与Gemma直接动作策略拥有相同任务、当前/目标图、公开访问历史和行动预算；属于不同策略比较，不是同策略下模型能力因果实验。')
    d.add_picture(str(FIG/'03_单模型同题工程对照.png'),width=Cm(16.4))
    table=d.add_table(rows=1,cols=5);table.style='Light Shading Accent 1'
    for cell,text in zip(table.rows[0].cells,['方案','完成/计划','成功/计划','SG米','API/回退']):cell.text=text
    for name,m in result['metrics'].items():
        values=[name,f"{m['completed']}/{m['planned']}",f"{m['successes']}/{m['planned']}",'-' if m['sg_m'] is None else f"{m['sg_m']:.1f}",f"{m['api_requests']}/{m['fallback_actions']}"]
        for cell,text in zip(table.add_row().cells,values):cell.text=text
    d.add_paragraph('Cloud错误时回退到本地提议；出现回退则应称混合执行成绩。连续三次错误后停止，并把未运行题保留在计划分母；未完整完成时不计算全队列SG。API联通或消息正确不证明稳定利用目标图，20题单地区也不支持地理泛化。')
    d.add_page_break();d.add_heading('中期交付验收与后续边界',1)
    for text in ['实时执行：真实本地推理、冻结轨迹回归及单步/自动/终止已分别核验。','数据与接口：API配置不回显凭据；图像限当前/目标或已访问格；正式备用区域没有消费。','证据：每次新运行写首动作前绑定、追加动作、终局及seal，旧实验/权重/阈值/验收未覆盖。','材料：方法图保存PNG/SVG/PDF及输入SHA，新Word和逐页渲染供中期报告/答辩使用，原用户稿不覆盖。']:
        d.add_paragraph(text,style='List Bullet')
    d.add_paragraph('当前主线有中期交付基础。新影像产品、15×15坐标/访问桶漂移、任意角度/异时/跨视角目标仍是后续研究，不能写成已完成能力；多角色协作原型有记录，但无足够收益证据升级默认。优先完成平台与材料，不为中期交付重开训练或扩大Agent数量。')
    d.add_heading('追溯入口',1)
    for text in ['项目当前状态：docs/PROJECT_STATE.md；任务和恢复：.agents/TASKS.md、HANDOFF.md。','方法原件：project/local_policy_default.json、project/defaults/masa_roads_grid10_radial_v1.json；绘图/平台方法图_v1/来源/。','本轮对照：DATA/processed_data/MassGIS/评测结果/平台单模型对照_v1；运行原件：DATA/processed_data/MassGIS/平台运行。','平台说明：平台/实时导航平台_v2/README.md；图表源数据与封存路径在Git收录说明中说明，Git克隆本身不含完整影像/权重/逐步日志。']:
        d.add_paragraph(text)
    footer=sec.footer.paragraphs[0];footer.alignment=2
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    output=REPORT/'主动探索视觉地理定位中期平台交付总结_v1.docx'
    if output.exists():raise ValueError('新报告已存在，不覆盖')
    d.save(output)
    (REPORT/'图表与方法说明_v1.md').write_text('# 中期平台方法与图表\n\n新报告不覆盖原用户Word。图1说明训练和实时闭环；图2给出cue/覆盖/径向参数与公式；图3为本轮20题单地区工程对照，非泛化结论。\n\n图源：绘图/平台方法图_v1；SHA在来源目录。原S2/S3/S4分别有效，MassGIS新域未过仍保留。平台实时推理复用编码缓存；系统有训练来源，当前默认单导航Agent。\n',encoding='utf-8')
    print(output)


if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('phase',choices=['methods','results','manifest','document']);phase=parser.parse_args().phase
    {'methods':method_figures,'results':result_figure,'manifest':source_manifest,'document':document}[phase]()
