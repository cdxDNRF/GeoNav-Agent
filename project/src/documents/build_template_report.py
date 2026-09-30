"""按课程模板重排选题报告：以模板解析副本为底，填入校正版内容。

- 基底：选题报告相关/模板解析工作区/综合工程设计-1-选题报告_模板解析.docx
  （由课程模板 .doc 用 Word COM 只读转换而来，模板原件不动）
- 内容源：选题报告相关/选题报告_..._校正版.docx（逐段按前缀抽取，保证文字保真）
- 输出：选题报告相关/选题报告_..._模板版.docx
- 模板栏目编号保持原样：一、二、三（自动编号）、四、六、七
"""
from pathlib import Path
import re
import shutil
import sys
import zipfile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

ROOT = Path(__file__).resolve().parents[3]
DOCS = ROOT / "选题报告相关"
WORK = DOCS / "模板解析工作区"
BASE = WORK / "综合工程设计-1-选题报告_模板解析.docx"
SRC = DOCS / "选题报告_基于主动探索的视觉地理定位方法设计与实现_校正版.docx"
OUT = DOCS / "选题报告_基于主动探索的视觉地理定位方法设计与实现_模板版.docx"
FIG = WORK / "fig_tech_route_from_corrected.png"

TITLE = "基于主动探索的视觉地理定位方法设计与实现"


# ---------- 读取内容源 ----------
def extract_source():
    with zipfile.ZipFile(SRC) as z:
        xml = z.read("word/document.xml").decode("utf-8")
        media = {n.rsplit("/", 1)[-1]: z.read(n) for n in z.namelist() if n.startswith("word/media/")}
    body = re.sub(r"<w:tbl>.*?</w:tbl>", "\u0001T\u0001", xml, flags=re.S)
    paras = []
    for seg in re.split(r"</w:p>", body):
        t = "".join(re.findall(r"<w:t[^>]*>([^<]*)</w:t>", seg)).strip()
        if t:
            paras.append(t)
    tables = []
    for tbl in re.findall(r"<w:tbl>.*?</w:tbl>", xml, re.S):
        grid = []
        for r in re.findall(r"<w:tr[ >].*?</w:tr>", tbl, re.S):
            row = []
            for c in re.findall(r"<w:tc>.*?</w:tc>", r, re.S):
                row.append(" ".join("".join(re.findall(r"<w:t[^>]*>([^<]*)</w:t>", c)).split()))
            grid.append(row)
        tables.append(grid)
    if len(media) != 1:
        raise ValueError(f"校正版图片数量异常: {list(media)}")
    FIG.write_bytes(next(iter(media.values())))
    return paras, tables


def para(paras, prefix):
    """按前缀取整段原文。"""
    for t in paras:
        if t.startswith(prefix):
            return t
    raise KeyError(prefix)


def rng(paras, start, end=None):
    i = next(k for k, t in enumerate(paras) if t.startswith(start))
    j = len(paras) if end is None else next(k for k, t in enumerate(paras) if t.startswith(end))
    return paras[i:j]


def sentence_with(paras, prefix, marker):
    """取某段中包含 marker 的完整句子（按句号切分）。"""
    text = para(paras, prefix)
    for s in re.split(r"(?<=。)", text):
        if marker in s:
            return s.strip()
    raise KeyError(marker)


# ---------- 写文档辅助 ----------
def set_run(run, size=12, bold=False, cn="宋体"):
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = None
    run._element.rPr.rFonts.set(qn("w:eastAsia"), cn)


def add_para(cell, text, kind="p", size=12):
    p = cell.add_paragraph()
    if kind == "h":
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.3
        p.paragraph_format.keep_with_next = True
        set_run(p.add_run(text), size=size, bold=True, cn="黑体")
    elif kind == "cap":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.3
        p.paragraph_format.keep_with_next = True
        set_run(p.add_run(text), size=10.5, bold=True)
    else:
        p.paragraph_format.line_spacing = 1.3
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.first_line_indent = Pt(24)
        set_run(p.add_run(text), size=size)
    return p


def add_figure(cell, path, width_cm=14.0):
    p = cell.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(0)
    set_run(p.add_run(), size=12)
    p.runs[0].add_picture(str(path), width=Cm(width_cm))
    return p


def set_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "000000")
        borders.append(el)
    tbl_pr.append(borders)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl_pr.append(layout)


def add_table(cell, grid, widths_cm):
    table = cell.add_table(rows=len(grid), cols=len(grid[0]))
    table.autofit = False
    set_borders(table)
    # 表头跨页重复、行内容不拆断（trPr 内元素顺序：cantSplit 在前，tblHeader 在后）
    for row in table.rows:
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
    header_trpr = table.rows[0]._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    header_trpr.append(header)
    for ri, row in enumerate(grid):
        for ci, text in enumerate(row):
            tc = table.cell(ri, ci)
            tc.width = Cm(widths_cm[ci])
            p = tc.paragraphs[0]
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            if ri == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                set_run(p.add_run(text), size=10.5, bold=True)
            else:
                set_run(p.add_run(text), size=10.5)
    spacer = cell.add_paragraph()
    spacer.paragraph_format.space_after = Pt(0)
    spacer.paragraph_format.line_spacing = 1.0
    set_run(spacer.add_run(""), size=6)
    return table


def clear_cell(cell):
    ps = cell.paragraphs
    for p in ps[1:]:
        p._element.getparent().remove(p._element)
    return cell


def rewrite_paragraph_first_run(p, text):
    runs = p.runs
    if not runs:
        p.add_run(text)
        return
    runs[0].text = text
    for r in runs[1:]:
        r._element.getparent().remove(r._element)


# ---------- 装配 ----------
def build():
    if OUT.exists():
        raise FileExistsError(f"输出已存在，拒绝覆盖：{OUT}")
    paras, tables = extract_source()
    shutil.copyfile(BASE, WORK / "_build_tmp.docx")
    doc = Document(WORK / "_build_tmp.docx")
    big = doc.tables[0]

    # 封面文字行
    for p in doc.paragraphs:
        t = p.text
        if re.match(r"^题\s+目：", t):
            rewrite_paragraph_first_run(p, f"题    目：{TITLE}")
            # 标题较长：14pt + 2字符缩进，避免末字单独换行
            p.runs[0].font.size = Pt(14)
            p.paragraph_format.first_line_indent = Pt(30)
            ind = p._element.pPr.find(qn("w:ind"))
            if ind is not None and qn("w:firstLineChars") in ind.attrib:
                del ind.attrib[qn("w:firstLineChars")]
        elif re.match(r"^填写时间：", t):
            rewrite_paragraph_first_run(p, "填写时间：2026 年 9 月")
    # 首页表格：综合工程设计题目
    first_cells = big.rows[0].cells
    target = first_cells[-1]
    clear_cell(target)
    p = target.paragraphs[0]
    rewrite_paragraph_first_run(p, TITLE)

    rows = {name: big.rows[i].cells[0] for i, name in
            zip(range(6, 12), ("一", "二", "三", "四", "六", "七"))}

    # ---- 一、项目组成员情况介绍 ----
    c = clear_cell(rows["一"])
    add_para(c, "本课题由三名大四学生组成的小组承担（计算机科学与技术学院，大数据智能方向）；成员姓名、学号与联系方式见封面登记表，成员分工见“成员分工情况”栏。")
    add_para(c, "知识条件：具备机器学习、计算机视觉、算法设计与分析等课程基础，熟悉 Python 编程与深度学习框架 PyTorch，了解大语言模型与视觉语言模型的 API 调用与提示工程方法；已共同精读 GOMAA-Geo、GeoExplorer 与 DynCur-Geo 三篇主动地理定位核心论文，掌握该任务的问题定义、统一评测协议（SR/SG）与现有方法的基本工作方式；已复现主动地理定位数据处理流程并完成实验平台的初步搭建。")
    add_para(c, "特长与兴趣：小组对视觉语言模型应用、智能体（Agent）系统设计与强化学习方向有浓厚兴趣，具备从数据管线、环境接口到评测脚本的工程实现能力，能够在零训练约束下组织实验与消融分析。")
    add_para(c, "科技活动：围绕本课题开展每周组会研讨与论文精读，按教学日历分阶段推进平台建设、机制实验与系统评测；相关分工与里程碑见表3与成员分工表。")

    # ---- 二、任务与要求 ----
    c = clear_cell(rows["二"])
    for t in rng(paras, "在地震搜救", "二、国内外研究现状"):
        add_para(c, t)
    add_para(c, "任务：", kind="h")
    for t in rng(paras, "总体目标：", "四、研究内容"):
        add_para(c, t)
    add_para(c, "要求：", kind="h")
    reqs = [
        sentence_with(paras, "总体目标：", "不进行本课题任务相关"),
        sentence_with(paras, "所有用于提升声明的方法", "同一冻结测试清单"),
        sentence_with(paras, "拟基于公开数据构建统一环境", "不得获取goal坐标"),
        sentence_with(paras, "SR与SG见相关AGL论文", "探索性假设"),
        para(paras, "成果目标：").replace("成果目标：", ""),
    ]
    for i, t in enumerate(reqs, 1):
        add_para(c, f"（{i}）{t}")

    # ---- 三、拟采用的研究方案和要解决的关键技术问题 ----
    c = clear_cell(rows["三"])
    add_para(c, "国内外研究现状", kind="h")
    for t in rng(paras, "（一）被动式视觉地理定位", "表 1"):
        add_para(c, t, kind="h" if t.startswith("（") else "p")
    for t in rng(paras, "表 1", "（三）研究空白"):
        if "表" in t and "对比" in t:
            add_para(c, t, kind="cap")
    add_table(c, tables[1], [2.3, 2.1, 3.3, 3.5, 3.2])
    for t in rng(paras, "（三）研究空白", "三、研究目标"):
        add_para(c, t, kind="h" if t.startswith("（") else "p")
    add_para(c, "研究内容与拟解决的关键问题", kind="h")
    for t in rng(paras, "（一）关键问题分析", "五、技术路线"):
        add_para(c, t, kind="h" if t.startswith("（") else "p")
    add_para(c, "技术路线与总体方案", kind="h")
    add_para(c, para(paras, "系统技术路线如图1所示"))
    add_figure(c, FIG)
    add_para(c, "图 1  系统总体技术路线图", kind="cap")
    for t in rng(paras, "其中，单个任务的决策闭环", "六、可行性分析"):
        add_para(c, t)

    # ---- 四、项目实施方案及实施计划 ----
    c = clear_cell(rows["四"])
    add_para(c, "实施方案", kind="h")
    add_para(c, "拟按六个阶段推进：先搭建数据管线与网格化环境、完成 SR/SG 评测与基线；再接入视觉语言模型、构建多智能体基础框架；随后开展记忆与层级搜索的机制实验；最后进行同协议对比、失败分析与报告撰写（里程碑见表3）。")
    add_para(c, para(paras, "主要风险与对策"))
    add_para(c, "可行性分析", kind="h")
    for t in rng(paras, "（一）技术可行性", "七、创新点"):
        add_para(c, t, kind="h" if t.startswith("（") else "p")
    add_para(c, "实施计划", kind="h")
    add_para(c, para(paras, "本课题按教学日历"))
    add_para(c, para(paras, "表 3"), kind="cap")
    add_table(c, tables[3], [3.4, 1.4, 2.2, 5.4, 2.0])

    # ---- 六、成员分工情况 ----
    c = clear_cell(rows["六"])
    add_para(c, para(paras, "表 4"), kind="cap")
    add_table(c, tables[4], [3.2, 3.6, 7.6])
    add_para(c, para(paras, "三人共同参与文献调研"))

    # ---- 七、预期成果及成果形式 ----
    c = clear_cell(rows["七"])
    rewrite_paragraph_first_run(c.paragraphs[0], "七、预期成果及成果形式")
    add_para(c, "预期成果：", kind="h")
    for i, t in enumerate(rng(paras, "主动地理定位网格化交互实验平台一套", "表 2"), 1):
        add_para(c, f"（{i}）{t}")
    add_para(c, para(paras, "表 2"), kind="cap")
    add_table(c, tables[2], [3.0, 7.0, 4.4])
    add_para(c, para(paras, "SR与SG见相关AGL论文"))
    add_para(c, "成果形式：", kind="h")
    add_para(c, "（1）网格化交互实验平台与方法源代码（含数据管线、环境接口、评测与轨迹可视化脚本）；（2）基于多智能体的主动定位方法实现（含 episode 内记忆与层级搜索模块）；（3）对比与消融实验的分析结果与图表；（4）综合工程设计报告。")

    # ---- 参考文献（表格之后） ----
    # 删除表后多余空段（模板遗留），避免连续空段检查
    body = doc.element.body
    tail_ps = [ch for ch in body.iterchildren(qn("w:p"))]
    for p_el in tail_ps[-2:]:
        if not "".join(t.text or "" for t in p_el.iter(qn("w:t"))).strip():
            body.remove(p_el)
    add_para_doc = doc.add_paragraph
    p = add_para_doc()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.line_spacing = 1.3
    set_run(p.add_run("参考文献"), size=14, bold=True, cn="黑体")
    for t in rng(paras, "[1] Sarkar", None):
        q = add_para_doc()
        q.paragraph_format.line_spacing = 1.3
        q.paragraph_format.space_after = Pt(0)
        set_run(q.add_run(t), size=10.5)

    doc.save(OUT)
    (WORK / "_build_tmp.docx").unlink(missing_ok=True)
    return OUT


if __name__ == "__main__":
    path = build()
    print(f"built: {path}")